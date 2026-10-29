"""Execute source → knowledge → PostgreSQL → chat with deterministic models.

Review judgments are fixture inputs, not approval of a real corpus. All commands
run in isolated temporary directories; PostgreSQL uses the disposable test schema.
"""
import copy
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys

from docx import Document
import pytest

from knowledge_foundry.retrieval.artifacts import load_artifacts, project
from knowledge_foundry.retrieval.config import ROOT, load_config
from knowledge_foundry.retrieval.provenance import graph_path
from knowledge_foundry.structural_chunks import assemble_structural_chunks, stable_hash
from claim_fixtures import fixture_settings


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2))
    return path


def command(script, *args, expected=0, env=None):
    result = subprocess.run([sys.executable, '-I', str(ROOT / 'scripts' / script), *map(str, args)],
                            cwd=ROOT, env={**os.environ, 'OPENAI_API_KEY': '', **(env or {})},
                            capture_output=True, text=True, timeout=60)
    assert result.returncode == expected, f'{script}: {result.stdout}\n{result.stderr}'
    return json.loads(result.stdout) if result.stdout.strip() else None


@pytest.fixture
def parsed_sources(tmp_path):
    (tmp_path / 'equipment.md').write_text('# Equipment\n\nOperators must inspect the equipment daily.\n')
    doc = Document()
    doc.add_heading('Records', level=1)
    doc.add_paragraph('Operators must retain inspection records for five years.')
    doc.save(tmp_path / 'records.docx')
    (tmp_path / 'exceptions.html').write_text(
        '<html><body><main><h1>Exceptions</h1><p>Emergency equipment requires an additional inspection.</p></main></body></html>')
    sources = []
    for name, media in [('equipment.md', 'text/markdown'),
                        ('records.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
                        ('exceptions.html', 'text/html')]:
        data = (tmp_path / name).read_bytes()
        sid = Path(name).stem
        sources.append({'manifest_id': sid, 'source_id': 'src:' + sid, 'title': sid.title(),
                        'class': 'test-policy', 'source_version': 'v1', 'publisher_url': 'https://example.org/' + name,
                        'local_path': name, 'retrieved_on': '2026-10-22', 'media_type': media,
                        'size_bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest()})
    manifest = write_json(tmp_path / 'manifest.json', {'schema_version': '1.0.0', 'sources': sources})
    config = write_json(tmp_path / 'pipeline.json', {
        'root': str(tmp_path), 'manifest_file': manifest.name, 'manifest_ids': [s['manifest_id'] for s in sources],
        'chunking': {'max_tokens': 256}, 'evaluation': {'chunking': {'sample_per_document': 12}},
        'structuring': {'provider': 'internal', 'options': {'html': {'content_selector': 'main'}, 'ocr': {'mode': 'off'}}}})
    assert all(o['status'] == 'ingested' for o in command('ingest.py', '--config', config, '--output-dir', tmp_path / 'artifacts/ingested')['outcomes'])
    parsed = tmp_path / 'parsed'
    assert all(o['status'] == 'parsed' for o in command('parse.py', '--config', config, '--output-dir', parsed)['outcomes'])
    return tmp_path, manifest, config, parsed


@pytest.mark.parametrize('format_name', ['markdown', 'docx'])
def test_parsing_evaluation_checks_native_locations(parsed_sources, format_name):
    root, _, _, parsed = parsed_sources
    report = command('deterministic_document_structuring_review.py', 'prepare', '--checkpoint', 'parsing', '--audit-only', '--parsed-dir', parsed, '--source-root', root)
    assert report['gate']['status'] == 'pass'
    bundle_path = next(p for p in parsed.glob('documents/*/artifact-bundle.json')
                       if json.loads(p.read_text())['sources'][0]['format'] == format_name)
    bundle = json.loads(bundle_path.read_text())
    location = bundle['elements'][-1]['location']
    location.update({'line_end': 99999} if format_name == 'markdown' else {'block_index': 99999})
    write_json(bundle_path, bundle)
    rejected = command('deterministic_document_structuring_review.py', 'prepare', '--checkpoint', 'parsing', '--audit-only', '--parsed-dir', parsed, '--source-root', root, expected=1)
    assert rejected['gate']['status'] == 'fail'
    assert rejected['corpus']['error_codes']['INVALID_LOCATION'] >= 1


def fixture_review(path):
    review = json.loads(path.read_text())
    review['reviewer'] = {'kind': 'agent', 'name': 'Deterministic fixture reviewer', 'model': None}
    for row in review['reviews']:
        row.update(status='pass', severity=None, notes='Fixture judgment for command/gate mechanics only.')
    write_json(path, review)


def test_documented_pipeline_produces_all_rag_inputs_and_vectors(parsed_sources, store, monkeypatch):
    root, manifest, pipeline_config, parsed = parsed_sources
    assert command('deterministic_document_structuring_review.py', 'prepare', '--checkpoint', 'parsing', '--audit-only', '--parsed-dir', parsed, '--source-root', root)['gate']['status'] == 'pass'
    chunked, context, generated, classified, summaries = [root / n for n in ['chunked', 'context', 'generated', 'classified', 'summaries']]
    command('chunk.py', '--config', pipeline_config, '--input-dir', parsed, '--output-dir', chunked)
    chunk_review_dir = root / 'chunk-review'
    chunk_args = ['--checkpoint', 'chunking', '--chunked-dir', chunked, '--parsed-dir', parsed, '--manifest', manifest, '--chunk-review-dir', chunk_review_dir]
    command('deterministic_document_structuring_review.py', 'prepare', *chunk_args)
    pending = json.loads((chunk_review_dir / 'report.json').read_text())
    assert pending['automatic_gate'] == 'pass' and pending['gate'] == 'pending_review'
    fixture_review(chunk_review_dir / 'review.json')
    command('deterministic_document_structuring_review.py', 'finalize', *chunk_args)
    assert json.loads((chunk_review_dir / 'report.json').read_text())['gate'] == 'pass'

    settings = fixture_settings()
    settings['steps']['embeddings'].update(dimensions=3, expected_dimensions=3)
    model_config = write_json(root / 'models.json', settings)
    bundle = json.loads((chunked / 'artifact-bundle.json').read_text())
    structural, _ = assemble_structural_chunks(bundle, json.loads((chunked / 'document-structure.json').read_text()))
    doc_ids = [d['id'] for d in bundle['documents']]
    vectors = {stable_hash(r['text']): [float(doc_ids.index(r['document_id']) == i) for i in range(3)] for r in structural}
    embedding_fixture = write_json(root / 'embeddings.json', {'embedding_vectors': vectors})
    from knowledge_foundry.document_structuring import accept_structure
    accept_structure(chunked, test_only=True)
    command('build_semantic_clusters.py', '--chunked-dir', chunked,
            '--output-dir', context, '--config', model_config, '--cache-dir', root / 'build-cache', '--fixture', embedding_fixture)
    clusters = json.loads((context / 'semantic-clusters.json').read_text())
    from knowledge_foundry.claim_generation_and_classification import _evidence_choices
    elements = {e['id']: e for e in bundle['elements']}
    chunks = {c['id']: c for c in bundle['chunks']}
    fixture = {'candidate_clusters': {}, 'classifications': {}}
    for cluster in clusters:
        choices = _evidence_choices([chunks[i] for i in cluster['original_chunk_ids']], elements)
        claims, coverage = [], []
        for label, passage in choices.items():
            substantive = passage['kind'] != 'heading'
            indices = [len(claims)] if substantive else []
            if substantive:
                claims.append({'text': passage['text'], 'evidence': [label]})
                fixture['classifications'][passage['text']] = 'topic:operations'
            coverage.append({'evidence_id': label, 'status': 'represented' if substantive else 'non_substantive',
                             'claim_indices': indices, 'reason': 'Deterministic fixture disposition.'})
        fixture['candidate_clusters'][','.join(cluster['original_chunk_ids'])] = {'claims': claims, 'coverage': coverage}
    claim_fixture = write_json(root / 'claims.json', fixture)
    taxonomy = write_json(root / 'taxonomy.json', {'taxonomy_version': 'test-v1', 'topics': [
        {'id': 'topic:operations', 'title': 'Operations', 'scope': 'Equipment inspection and record retention.'}]})
    command('generate_claims.py', '--context-dir', context, '--output-dir', generated, '--config', model_config, '--fixture', claim_fixture)
    command('classify_claims.py', '--candidates-dir', generated, '--output-dir', classified,
            '--config', model_config, '--taxonomy', taxonomy, '--fixture', claim_fixture)
    review_args = ['--candidates-dir', generated, '--classified-dir', classified, '--manifest', manifest,
                   '--source-root', root, '--taxonomy', taxonomy]
    assert command('deterministic_claim_review.py', 'prepare', *review_args)['gate'] == 'pending_review'
    fixture_review(classified / 'compilation-review.json')
    assert command('deterministic_claim_review.py', 'finalize', *review_args, '--review-file', classified / 'compilation-review.json')['gate'] == 'pass'

    summary_config = json.loads((ROOT / 'config/summaries.json').read_text())
    summary_config['steps']['summary_generation'].update(provider='fixture', retries=0)
    summary_config['steps']['embeddings'].update(provider='fixture', dimensions=3, expected_dimensions=3)
    summary_config['pricing']['providers']['fixture'] = copy.deepcopy(summary_config['pricing']['providers']['openai'])
    summary_config['steps']['summary_clusters']['max_levels'] = 1
    summary_cfg = write_json(root / 'summary-config.json', summary_config)
    leaf_texts = {c['id']: 'Policy summary: ' + ' '.join(chunks[cid]['text'] for cid in c['original_chunk_ids'])
                  for c in clusters}
    summary_fixture = {'summaries': {cid: {'text': text} for cid, text in leaf_texts.items()},
                       'embedding_vectors': {stable_hash(text): [1., 0., 0.] for text in leaf_texts.values()}}
    summary_fixture_path = write_json(root / 'summary-fixture.json', summary_fixture)
    # Discover content-addressed leaf IDs to supply a deterministic higher-level response.
    preview = root / 'summary-fixture-plan'
    summary_args = ['--context-dir', context, '--config', summary_cfg, '--cache-dir', root / 'summary-cache', '--fixture', summary_fixture_path]
    command('generate_summaries.py', *summary_args, '--output-dir', preview)
    leaf_ids = sorted(json.loads(line)['id'] for line in (preview / 'summaries.jsonl').read_text().splitlines())
    group_id = 'summary-cluster:' + stable_hash(['recursive-summaries-v1', 2, leaf_ids])
    summary_fixture['summaries'][group_id] = {'text': 'Inspection and record-retention policies across the documents.'}
    write_json(summary_fixture_path, summary_fixture)
    summary_config['steps']['summary_clusters']['max_levels'] = 2
    write_json(summary_cfg, summary_config)
    command('generate_summaries.py', *summary_args, '--output-dir', summaries)
    graph = classified / 'knowledge-provenance-graph.html'
    command('render_knowledge_graph.py', '--bundle', classified / 'artifact-bundle.json',
            '--claim-generation-dir', generated, '--summary-dir', summaries, '--output', graph)

    cfg = load_config()
    cfg['search'].update(language='english', method='keyword', rerank=False)
    cfg['answer']['provider'] = 'extractive'
    cfg['embedding'].update(provider='fixture', dimensions=3, expected_dimensions=3)
    cfg['cache_dir'] = str(root / 'rag-cache')
    retrieval_config = write_json(root / 'retrieval.json', cfg)
    loaded = load_artifacts(preparation_dir=context, claims_dir=classified, summaries_dir=summaries)
    records, _ = project(loaded, cfg['embedding'])
    assert any(r['kind'] == 'claim' and r['vector'] is None for r in records)
    assert any(r['kind'] == 'summary' and r['vector'] is None for r in records)
    assert any(r['kind'] == 'summary' and r['vector'] is not None for r in records)
    fixture_path = write_json(root / 'rag-vectors.json', {'embedding_vectors': {stable_hash(r['text']): [1., 0., 0.] for r in records}})
    frozen = {p: hashlib.sha256(p.read_bytes()).hexdigest() for directory in [context, generated, classified, summaries]
              for p in directory.rglob('*') if p.is_file()}
    env = {'KF_DATABASE_URL': store.dsn}
    command('index_artifacts.py', 'migrate', env=env)
    import_args = ['import', '--corpus', 'equipment', '--config', retrieval_config, '--preparation-dir', context,
                   '--claims-dir', classified, '--summaries-dir', summaries, '--embed-missing', '--fixture', fixture_path]
    first = command('index_artifacts.py', *import_args, env=env)
    second = command('index_artifacts.py', *import_args, env=env)
    assert first['created'] and not second['created'] and first['snapshot_id'] == second['snapshot_id']
    with pytest.raises(ValueError, match='published'):
        store.snapshot('equipment')
    command('index_artifacts.py', 'publish', '--corpus', 'equipment', '--snapshot-id', first['snapshot_id'], env=env)
    assert all(c['records'] > 0 and c['records'] == c['vectors'] for c in first['vector_coverage'].values())
    snapshot = store.snapshot('equipment')
    definition = {'title': 'Equipment', 'runs': {'test': {'claims_dir': str(classified), 'graph_path': str(graph)}}}
    assert graph_path(snapshot, definition) == graph
    question = 'Operators must inspect the equipment daily.'
    for mode in ['passage', 'claim', 'summary']:
        answer = command('query.py', question, '--corpus', 'equipment', '--config', retrieval_config,
                         '--method', 'keyword', '--mode', mode, '--provider', 'extractive', env=env)
        assert answer['status'] == 'extractive', (mode, answer['status'], answer.get('retrieval', {}).get('warnings'))
        assert answer['citation_evaluation']['citation_resolution_rate'] == 1
    cases = write_json(root / 'cases.json', [{'id': 'inspection', 'question': question, 'source_id': 'src:equipment',
                                             'text_anchors': ['inspect the equipment daily']}])
    evaluation = root / 'rag-evaluation.json'
    evaluation_args = ['--corpus', 'equipment', '--config', retrieval_config, '--cases', cases,
                       '--methods', 'keyword', '--modes', 'passage', 'claim', 'summary', '--no-rerank', '--check-excerpts', '--output', evaluation]
    command('evaluate_retrieval.py', *evaluation_args, env=env)
    assert all(m['text_anchor_hits'] == 1 for m in json.loads(evaluation.read_text())['metrics'])
    write_json(cases, [{'id': 'oversized', 'question': 'word ' * 2000, 'source_id': 'src:equipment'}])
    command('evaluate_retrieval.py', *evaluation_args, env=env, expected=1)
    assert all('error' in r for r in json.loads(evaluation.read_text())['rows'])
    assert all(hashlib.sha256(p.read_bytes()).hexdigest() == digest for p, digest in frozen.items())

    from streamlit.testing.v1 import AppTest
    monkeypatch.setenv('KF_DATABASE_URL', store.dsn)
    monkeypatch.setenv('KF_RETRIEVAL_CONFIG', str(retrieval_config))
    monkeypatch.setenv('KF_CORPORA_CONFIG', str(write_json(root / 'corpora.json', {'equipment': definition})))
    app = AppTest.from_file(str(ROOT / 'apps/chat/app.py'), default_timeout=20).run()
    assert not app.exception and any(b.label == 'Open provenance graph' for b in app.button)
    app.chat_input[0].set_value(question).run()
    assert not app.exception and any('inspect the equipment daily' in c.value for c in app.code)


def test_parsing_review_cli_preserves_pending_and_rejects_stale_inputs(parsed_sources):
    root, manifest, _, parsed = parsed_sources
    folder = root / 'parsing-review'
    review_path, report_path = folder / 'review.json', folder / 'report.json'
    args = ['--checkpoint', 'parsing', '--parsed-dir', parsed, '--source-root', root, '--manifest', manifest, '--parsing-review-dir', folder]
    command('deterministic_document_structuring_review.py', 'prepare', *args)
    pending = json.loads(report_path.read_text())
    assert pending['automatic_gate'] == 'pass' and pending['gate'] == 'pending_review'
    original = review_path.read_bytes()
    command('deterministic_document_structuring_review.py', 'prepare', *args)
    assert review_path.read_bytes() == original
    command('deterministic_document_structuring_review.py', 'finalize', *args, expected=1)
    review = json.loads(original)
    review['reviewer'] = {'kind': 'agent', 'name': 'Fixture only', 'model': 'fixture'}
    for row in review['reviews']:
        row.update(status='pass', notes='Fixture judgment never real approval',
                   inspection={'original_inspected': True, 'layout_inspected': False})
    write_json(review_path, review)
    command('deterministic_document_structuring_review.py', 'finalize', *args)
    assert json.loads(report_path.read_text())['gate'] == 'pass'
    (root / 'equipment.md').write_text('changed acquisition path; frozen bytes remain valid')
    command('deterministic_document_structuring_review.py', 'finalize', *args)
    metadata = next(json.loads(p.read_text()) for p in parsed.glob('documents/*/parse.json')
                    if json.loads(p.read_text())['parser'] == 'python-docx')
    Path(metadata['ingested_artifact_file']).write_bytes(b'tampered ingested artifact')
    command('deterministic_document_structuring_review.py', 'finalize', *args, expected=2)
