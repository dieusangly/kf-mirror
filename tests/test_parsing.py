import hashlib
import io
import json
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from docx import Document
from jsonschema import Draft202012Validator, FormatChecker
from lxml import html

from knowledge_foundry.parsing import PARSERS, SCHEMA, _save_pdf_figure_assets, main, parse_collection
from knowledge_foundry.parsers import (
    _pdf_running_headers, _pdf_table_cells, _pdf_table_text, _pdf_text, parse_docx, parse_markdown, parse_pdf,
)


ROOT = Path(__file__).resolve().parents[1]


def _fixture(root, name, media_type, contents):
    path = root / name
    if isinstance(contents, bytes):
        path.write_bytes(contents)
    else:
        path.write_text(contents, encoding="utf-8")
    raw = path.read_bytes()
    return {"manifest_id": name, "source_id": f"src:{name}", "title": name, "class": "test",
            "source_version": "v1", "publisher_url": f"https://example.org/{name}",
            "local_path": name, "retrieved_on": "2026-10-07", "media_type": media_type,
            "size_bytes": len(raw), "sha256": hashlib.sha256(raw).hexdigest()}


def _config(root, entries):
    (root / "manifest.json").write_text(json.dumps({"schema_version": "1.0.0", "sources": entries}))
    data = {"root": ".", "manifest_file": "manifest.json",
            "manifest_ids": [entry["manifest_id"] for entry in entries],
            "parsing": {"ocr": {"mode": "off"}}}
    path = root / "config.json"
    path.write_text(json.dumps(data))
    return path


class ParsingTests(unittest.TestCase):
    def test_discretionary_hyphen_space_joins_source_word(self):
        self.assertEqual(_pdf_text('electronic instru\u00ad ments'), 'electronic instruments')

    def test_spaced_legal_headings_and_chapter_titles_reset_peer_scope(self):
        from knowledge_foundry.parsers import _pdf_heading_level
        stack = [(1, 'chapter', 'CHAPTER II', 'legal'), (3, 'section', 'Section 1', 'legal'),
                 (5, 'article', 'Article 8', 'legal')]
        self.assertEqual(_pdf_heading_level('S e c t i o n 3', stack, {}), (3, 'legal'))
        self.assertEqual(_pdf_heading_level('Chapitre 5 - A new chapter', stack, {}), (1, 'legal'))

    def test_english_and_standalone_annexes_reset_heading_scope(self):
        from knowledge_foundry.parsers import _pdf_heading_level
        stack = [(1, 'old', 'IV. The role of supervisors', 'numbered')]
        for heading in ('Annex 1', 'Annexe 2', 'ANNEXE'):
            self.assertEqual(_pdf_heading_level(heading, stack, {}), (1, 'numbered'))

    def test_source_verified_short_picture_transcript_is_retained(self):
        from docling_core.types.doc import (DoclingDocument, DocItemLabel, Size, BoundingBox,
                                           CoordOrigin, ProvenanceItem)
        doc = DoclingDocument(name='Fixture'); doc.add_page(1, Size(width=600, height=800))
        picture = doc.add_picture(prov=ProvenanceItem(page_no=1, charspan=(0, 0),
            bbox=BoundingBox(l=40, t=100, r=500, b=400, coord_origin=CoordOrigin.TOPLEFT)))
        child = doc.add_text(DocItemLabel.TEXT, 'Sensor', parent=picture, prov=ProvenanceItem(
            page_no=1, charspan=(0, 6), bbox=BoundingBox(l=50, t=120, r=100, b=130,
                                                                    coord_origin=CoordOrigin.TOPLEFT)))
        proof = {'repairs': [{'method': 'source_pixel_ocr', 'ref': child.self_ref}], 'uncertainties': []}
        converter = Mock(); converter.convert.return_value = SimpleNamespace(status='success', document=doc)
        with patch('knowledge_foundry.pdf_source.adapt_native', return_value=(doc.export_to_dict(), proof)):
            items, _, warnings = parse_pdf(Path('unused.pdf'), {}, {'ocr': {'mode': 'auto'}}, converter)
        self.assertTrue(any(item['text'] == 'Sensor' and item['details']['docling_ref'] == child.self_ref
                            for item in items))
        self.assertFalse(warnings)

    def test_native_rebuild_reverifies_identity_and_rebuilds_without_pdf_conversion(self):
        from docling_core.types.doc import (DoclingDocument, DocItemLabel, Size, BoundingBox,
                                           CoordOrigin, ProvenanceItem)
        from test_pdf_tables import write_pdf
        with tempfile.TemporaryDirectory() as name:
            root = Path(name)
            path = root / 'source.pdf'
            write_pdf(path, b'BT /F1 12 Tf 50 700 Td (Sensor reading.) Tj ET')
            entry = _fixture(root, path.name, 'application/pdf', path.read_bytes())
            config = _config(root, [entry])
            doc = DoclingDocument(name='Fixture'); doc.add_page(1, Size(width=600, height=800))
            doc.add_text(DocItemLabel.TEXT, 'Sensor reading.', prov=ProvenanceItem(page_no=1,
                charspan=(0, 15), bbox=BoundingBox(l=48, t=85, r=160, b=110,
                                                coord_origin=CoordOrigin.TOPLEFT)))
            converter = Mock(); converter.convert.return_value = SimpleNamespace(status='success', document=doc)
            with patch('knowledge_foundry.parsing.build_pdf_converter', return_value=converter):
                initial = parse_collection(config, root / 'prior')['outcomes'][0]
            prior = Path(initial['output'])
            saved_bundle = json.loads((prior / 'artifact-bundle.json').read_text())
            saved_parse = json.loads((prior / 'parse.json').read_text())
            with patch('knowledge_foundry.parsing.build_pdf_converter') as build:
                outcome = parse_collection(config, root / 'rebuilt', native_input_dir=root / 'prior')['outcomes'][0]
            build.assert_not_called()
            self.assertEqual(outcome['status'], 'parsed')
            rebuilt = Path(outcome['output'])
            self.assertEqual(json.loads((rebuilt / 'artifact-bundle.json').read_text()), saved_bundle)
            replay = json.loads((rebuilt / 'parse.json').read_text())['native_replay']
            self.assertEqual(replay['source_sha256'], entry['sha256'])
            self.assertEqual(replay['native_sha256'], hashlib.sha256((prior / 'docling.json').read_bytes()).hexdigest())
            for change in ['identity', 'settings', 'missing_native']:
                with self.subTest(change=change):
                    bundle = json.loads(json.dumps(saved_bundle))
                    settings = json.loads(json.dumps(saved_parse))
                    if change == 'identity':
                        bundle['documents'][0]['content_sha256'] = '0' * 64
                    elif change == 'settings':
                        settings['settings']['ocr']['mode'] = 'force'
                    else:
                        (prior / 'docling.json').unlink()
                    (prior / 'artifact-bundle.json').write_text(json.dumps(bundle))
                    (prior / 'parse.json').write_text(json.dumps(settings))
                    outcome = parse_collection(config, root / 'rejected', native_input_dir=root / 'prior')['outcomes'][0]
                    self.assertEqual(outcome['status'], 'failed')
            with self.assertRaisesRegex(ValueError, 'directories must differ'):
                parse_collection(config, root / 'prior', native_input_dir=root / 'prior')

    def test_uncertain_picture_transcript_is_retained_with_provenance(self):
        from docling_core.types.doc import (DoclingDocument, DocItemLabel, Size, BoundingBox,
                                           CoordOrigin, ProvenanceItem)
        doc = DoclingDocument(name='Fixture'); doc.add_page(1, Size(width=600, height=800))
        def prov(box, value):
            return ProvenanceItem(page_no=1, charspan=(0, len(value)), bbox=BoundingBox(
                l=box[0], t=box[1], r=box[2], b=box[3], coord_origin=CoordOrigin.TOPLEFT))
        picture = doc.add_picture(prov=prov([40, 100, 500, 400], ''))
        child = doc.add_text(DocItemLabel.TEXT, 'ZZ', parent=picture, prov=prov([50, 120, 80, 130], 'ZZ'))
        native = doc.export_to_dict()
        proof = {'uncertainties': [{'code': 'source_ocr_verification_uncertain', 'ref': child.self_ref,
                                   'page': 1, 'bbox': [50, 120, 80, 130]}]}
        converter = Mock(); converter.convert.return_value = SimpleNamespace(status='success', document=doc)
        with patch('knowledge_foundry.pdf_source.adapt_native', return_value=(native, proof)):
            items, original, warnings = parse_pdf(Path('unused.pdf'), {}, {'ocr': {'mode': 'auto'}}, converter)
        retained = next(item for item in items if item['text'] == 'ZZ')
        self.assertEqual(retained['details']['docling_ref'], child.self_ref)
        self.assertEqual(retained['location']['page'], 1)
        self.assertEqual(retained['details']['evidence_role'], 'auxiliary')
        self.assertEqual(original, native)
        self.assertEqual([w['code'] for w in warnings], ['source_ocr_verification_uncertain'])

    def test_short_numbered_native_picture_notes_preserve_document_evidence(self):
        values = [('footnote', '14 Keep the calibration certificate'),
                  ('caption', '15 Sensor wiring'), ('footnote', '16 sensor-manual.pdf')]
        picture = SimpleNamespace(label='picture', text='', marker=None, self_ref='#/pictures/0',
                                  parent=SimpleNamespace(cref='#/body'),
                                  prov=[SimpleNamespace(page_no=1, bbox=None)])
        nodes = [{'self_ref': f'#/texts/{i}', 'label': label, 'text': value,
                  'parent': {'$ref': picture.self_ref}, 'prov': [{'page_no': 1}]}
                 for i, (label, value) in enumerate(values)]
        native = {'texts': nodes, 'pictures': [{'self_ref': picture.self_ref,
                  'prov': [{'page_no': 1}], 'children': [{'$ref': n['self_ref']} for n in nodes]}],
                  'tables': [], 'groups': []}
        document = SimpleNamespace(iterate_items=lambda with_groups: iter([(picture, 1)]),
                                   pages={1: SimpleNamespace(size=SimpleNamespace(height=800))},
                                   export_to_dict=lambda: native)
        converter = Mock()
        converter.convert.return_value = SimpleNamespace(status='success', document=document)
        items, _, warnings = parse_pdf(Path('unused.pdf'), {}, {'ocr': {'mode': 'off'}}, converter)
        self.assertEqual(warnings, [])
        self.assertEqual([item['kind'] for item in items], ['figure', 'footnote', 'caption', 'footnote'])
        self.assertEqual([item['details']['evidence_role'] for item in items],
                         ['structural', 'authoritative', 'auxiliary', 'structural'])
        self.assertIsNone(items[1]['parent_key'])
        self.assertEqual(items[2]['parent_key'], items[0]['locator_key'])
        self.assertEqual(items[3]['details']['link_text_role'], 'non_substantive_filename')

    def test_picture_caption_yielded_again_by_outer_traversal_is_represented_once(self):
        caption = 'Bench arrangement for sensor testing.'
        picture = SimpleNamespace(label='picture', text='', marker=None, self_ref='#/pictures/0',
                                  parent=SimpleNamespace(cref='#/body'),
                                  prov=[SimpleNamespace(page_no=1, bbox=None)])
        child = SimpleNamespace(label='caption', text=caption, marker=None, self_ref='#/texts/0',
                                parent=SimpleNamespace(cref='#/pictures/0'),
                                prov=[SimpleNamespace(page_no=1, bbox=None)])
        native = {'texts': [{'self_ref': child.self_ref, 'label': 'caption', 'text': caption,
                             'parent': {'$ref': picture.self_ref}, 'prov': [{'page_no': 1}]}],
                  'pictures': [{'self_ref': picture.self_ref, 'prov': [{'page_no': 1}],
                                'children': [{'$ref': child.self_ref}]}], 'tables': [], 'groups': []}
        document = SimpleNamespace(iterate_items=lambda with_groups: iter([(picture, 1), (child, 2)]),
                                   pages={1: SimpleNamespace(size=SimpleNamespace(height=800))},
                                   export_to_dict=lambda: native)
        converter = Mock()
        converter.convert.return_value = SimpleNamespace(status='success', document=document)
        items, _, warnings = parse_pdf(Path('unused.pdf'), {}, {'ocr': {'mode': 'off'}}, converter)
        self.assertEqual(warnings, [])
        self.assertEqual([item['kind'] for item in items], ['figure', 'caption'])
        self.assertEqual(items[1]['details']['evidence_role'], 'auxiliary')
        self.assertEqual(items[1]['parent_key'], items[0]['locator_key'])

    def test_explicit_run_identity_for_independent_markdown_document(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            entry = _fixture(root, "library.md", "text/markdown",
                             "# Library\n\nReturn borrowed books within fourteen days.\n")
            config = _config(root, [entry])
            report = parse_collection(config, root / "output", run_id="library-audit")
            self.assertEqual(report["run_id"], "library-audit")
            self.assertEqual(report["outcomes"][0]["status"], "parsed")
            self.assertEqual(json.loads((root / "output/report.json").read_text())["run_id"],
                             "library-audit")
            with self.assertRaisesRegex(ValueError, "safe directory name"):
                parse_collection(config, root / "invalid", run_id="../other")

    def test_parse_pdf_uses_supplied_converter(self):
        document = SimpleNamespace(iterate_items=lambda with_groups: iter(()), pages={},
                                   export_to_dict=lambda: {"texts": [], "pictures": [],
                                                          "tables": [], "groups": []})
        converter = Mock()
        converter.convert.return_value = SimpleNamespace(status="success", document=document)
        with patch("docling.document_converter.DocumentConverter") as factory:
            items, _native, warnings = parse_pdf(Path("unused.pdf"), {}, {"ocr": {"mode": "off"}},
                                                 converter=converter)
        factory.assert_not_called()
        converter.convert.assert_called_once_with(Path("unused.pdf"))
        self.assertEqual((items, warnings), ([], []))

    def test_pdf_table_serializes_cells_once_and_preserves_empty_columns(self):
        def cell(row, column, value):
            return SimpleNamespace(start_row_offset_idx=row,
                                   start_col_offset_idx=column, text=value)
        index = SimpleNamespace(data=SimpleNamespace(table_cells=[
            cell(0, 0, "INTRODUCTION"), cell(0, 2, "2"),
            cell(1, 0, "La mise en œuvre selon"),
            cell(2, 0, "les risques"), cell(2, 3, "6"),
            cell(3, 0, "Oui"), cell(3, 1, "Oui"),
        ]))
        self.assertEqual(_pdf_table_text(index, None),
                         "INTRODUCTION |  | 2\nLa mise en œuvre selon\nles risques |  |  | 6\nOui | Oui")

    def test_pdf_table_does_not_guess_merged_headers_from_french_words(self):
        def cell(row, column, value, header=False, left=None):
            return SimpleNamespace(start_row_offset_idx=row, start_col_offset_idx=column,
                                   text=value, row_span=1, col_span=1,
                                   column_header=header, row_header=False,
                                   bbox=SimpleNamespace(l=left) if left is not None else None)
        table = SimpleNamespace(data=SimpleNamespace(table_cells=[
            cell(0, 0, "Thématique", True, 0),
            cell(0, 1, "Les sociétés de domiciliation Les agents sportifs", True, 5),
            cell(0, 2, "Le secteur du jeu Critères concernant particulièrement les déclarants", True, 10),
            cell(1, 0, "Risque"), cell(1, 1, "x"), cell(1, 3, "x"),
        ]))
        self.assertEqual(_pdf_table_text(table, None),
                         "Thématique | Les sociétés de domiciliation Les agents sportifs | Le secteur du jeu Critères concernant particulièrement les déclarants\nRisque | x |  | x")
        headers = [cell for cell in _pdf_table_cells(table) if cell["row"] == 0]
        self.assertEqual([(cell["column"], cell["text"]) for cell in headers], [
            (0, "Thématique"), (1, "Les sociétés de domiciliation Les agents sportifs"),
            (2, "Le secteur du jeu Critères concernant particulièrement les déclarants"),
        ])

    def test_filename_only_pdf_footnote_is_structural_but_meaningful_note_remains_evidence(self):
        def node(value, ref):
            return SimpleNamespace(label="footnote", text=value, marker=None,
                                   prov=[SimpleNamespace(page_no=1, bbox=None)], self_ref=ref,
                                   parent=SimpleNamespace(cref="#/body"))
        filename = node("23 20220404_lignes_directrices_revisees_relatives_identification.pdf", "#/texts/0")
        note = node("24 Voir les lignes directrices ACPR sur l'identification.", "#/texts/1")
        document = SimpleNamespace(iterate_items=lambda with_groups: iter([(filename, 1), (note, 1)]),
                                   pages={1: SimpleNamespace(size=SimpleNamespace(height=800))},
                                   export_to_dict=lambda: {"texts": [], "pictures": [], "tables": [], "groups": []})
        with patch("docling.document_converter.DocumentConverter") as converter:
            converter.return_value.convert.return_value = SimpleNamespace(status="success", document=document)
            items, _, _ = parse_pdf(Path("unused.pdf"), {}, {"ocr": {"mode": "off"}})
        self.assertEqual([item["details"]["evidence_role"] for item in items],
                         ["structural", "authoritative"])
        self.assertEqual(items[0]["details"]["link_text_role"], "non_substantive_filename")

    def test_pdf_text_removes_controls_and_layout_word_break(self):
        self.assertEqual(_pdf_text("INTRODUCTION\b\x07 et malus écolo- gique"),
                         "INTRODUCTION et malus écologique")

    def test_pdf_table_row_boundaries_are_not_prose_word_breaks(self):
        values = [('Entry', 'BC-'), ('FT index', '79'), ('Wrapped word', 'écolo-\ngique'),
                  ('URL', 'http://www.currency-\niso.org/file.xls')]
        cells = [SimpleNamespace(text=text, start_row_offset_idx=row, start_col_offset_idx=col,
                                 row_span=1, col_span=1, column_header=row == 0, row_header=False, bbox=None)
                 for row, pair in enumerate(values) for col, text in enumerate(pair)]
        node = SimpleNamespace(label='table', self_ref='#/tables/0', parent=SimpleNamespace(cref='#/body'),
                               prov=[SimpleNamespace(page_no=1, bbox=None)],
                               data=SimpleNamespace(table_cells=cells))
        document = SimpleNamespace(pages={1: SimpleNamespace(size=SimpleNamespace(height=800, width=600))},
            iterate_items=lambda with_groups: iter([(node, 1)]),
            export_to_dict=lambda: {'pages': {'1': {'size': {'width': 600, 'height': 800}}},
                                    'texts': [], 'pictures': [], 'tables': [], 'groups': []})
        converter = Mock()
        converter.convert.return_value = SimpleNamespace(status='success', document=document)
        items, _, warnings = parse_pdf(Path('unused.pdf'), {}, {'ocr': {'mode': 'off'}}, converter)
        self.assertEqual(warnings, [])
        table = next(i for i in items if i['kind'] == 'table')
        self.assertEqual(table['text'],
            'Entry | BC-\nFT index | 79\nWrapped word | écologique\nURL | http://www.currency-iso.org/file.xls')

    def test_wrapped_urls_keep_literal_hyphens_while_prose_is_dehyphenated(self):
        self.assertEqual(_pdf_text('écolo- gique https://www.example- org.test/a- b.xls '
                                   'et www.other- host.test puis régu- lier'),
                         'écologique https://www.example-org.test/a-b.xls '
                         'et www.other-host.test puis régulier')

    def test_native_heading_followed_by_body_stays_authoritative_below_figure(self):
        from docling_core.types.doc import BoundingBox, CoordOrigin

        def node(label, text, ref, box):
            return SimpleNamespace(label=label, text=text, self_ref=ref, marker=None,
                parent=SimpleNamespace(cref='#/body'), prov=[SimpleNamespace(page_no=1,
                bbox=BoundingBox(l=box[0], t=box[1], r=box[2], b=box[3],
                                 coord_origin=CoordOrigin.TOPLEFT))])
        picture = node('picture', '', '#/pictures/0', [50, 100, 500, 400])
        heading = node('section_header', 'VII. Equipment maintenance', '#/texts/0', [50, 430, 250, 442])
        subtitle = node('section_header', 'Inspection records', '#/texts/1', [50, 450, 200, 462])
        body = node('text', 'Operators must record the date of every inspection and retain '
                    'the results with the equipment maintenance log.', '#/texts/2', [50, 480, 500, 515])
        native = {'pages': {'1': {'size': {'width': 600, 'height': 800}}},
                  'pictures': [{'self_ref': picture.self_ref, 'children': [],
                    'prov': [{'page_no': 1, 'bbox': picture.prov[0].bbox.model_dump(mode='json')}]}],
                  'texts': [], 'tables': [], 'groups': []}
        document = SimpleNamespace(iterate_items=lambda with_groups: iter(
            (item, 1) for item in [picture, heading, subtitle, body]),
            pages={1: SimpleNamespace(size=SimpleNamespace(height=800))},
            export_to_dict=lambda: native)
        converter = Mock()
        converter.convert.return_value = SimpleNamespace(status='success', document=document)
        items, _, _ = parse_pdf(Path('unused.pdf'), {}, {'ocr': {'mode': 'off'}}, converter)
        self.assertEqual([i['kind'] for i in items], ['figure', 'heading', 'heading', 'paragraph'])
        self.assertEqual(items[-1]['details']['heading_path'],
                         ['VII. Equipment maintenance', 'Inspection records'])
        self.assertEqual(items[1]['details']['evidence_role'], 'authoritative')

    def test_control_only_pdf_node_is_not_missing_content(self):
        node = SimpleNamespace(label="text", text="\b", marker=None,
                               prov=[SimpleNamespace(page_no=1, bbox=None)],
                               self_ref="#/texts/0", parent=SimpleNamespace(cref="#/body"))
        document = SimpleNamespace(iterate_items=lambda with_groups: iter([(node, 1)]),
                                   pages={1: SimpleNamespace(size=SimpleNamespace(height=800))},
                                   export_to_dict=lambda: {"texts": [{"self_ref": "#/texts/0",
                                       "text": "\b", "parent": {"$ref": "#/body"}}],
                                       "pictures": [], "tables": [], "groups": []})
        with patch("docling.document_converter.DocumentConverter") as converter:
            converter.return_value.convert.return_value = SimpleNamespace(
                status="success", document=document)
            items, _, warnings = parse_pdf(Path("unused.pdf"), {}, {"ocr": {"mode": "off"}})
        self.assertEqual(items, [])
        self.assertEqual(warnings, [])

    def test_page_edge_drop_cap_is_structural(self):
        node = SimpleNamespace(label="text", text="F", marker=None,
                               prov=[SimpleNamespace(page_no=2, bbox=SimpleNamespace(
                                   to_top_left_origin=lambda height: SimpleNamespace(
                                       l=35, t=60, r=70, b=85)))],
                               self_ref="#/texts/0", parent=SimpleNamespace(cref="#/body"))
        native = {"pages": {"2": {"size": {"width": 600, "height": 800}}},
                  "texts": [{"self_ref": "#/texts/0", "text": "F",
                             "parent": {"$ref": "#/body"}}], "pictures": [], "tables": [], "groups": []}
        document = SimpleNamespace(iterate_items=lambda with_groups: iter([(node, 1)]),
                                   pages={2: SimpleNamespace(size=SimpleNamespace(height=800))},
                                   export_to_dict=lambda: native)
        with patch("docling.document_converter.DocumentConverter") as converter:
            converter.return_value.convert.return_value = SimpleNamespace(
                status="success", document=document)
            items, _, warnings = parse_pdf(Path("unused.pdf"), {}, {"ocr": {"mode": "off"}})
        self.assertEqual(warnings, [])
        self.assertEqual(items[0]["details"]["evidence_role"], "structural")

    def test_sparse_page_short_ocr_fragments_are_structural(self):
        def node(value, ref, top):
            return SimpleNamespace(label="text", text=value, marker=None,
                                   prov=[SimpleNamespace(page_no=2, bbox=SimpleNamespace(
                                       to_top_left_origin=lambda height: SimpleNamespace(
                                           l=100, t=top, r=200, b=top + 20)))],
                                   self_ref=ref, parent=SimpleNamespace(cref="#/body"))
        title = node("Chapter title", "#/texts/0", 100)
        noise = node("EDP", "#/texts/1", 600)
        native = {"pages": {"2": {"size": {"width": 600, "height": 800}}},
                  "texts": [{"self_ref": item.self_ref, "text": item.text,
                             "prov": [{"page_no": 2}], "parent": {"$ref": "#/body"}}
                            for item in (title, noise)],
                  "pictures": [], "tables": [], "groups": []}
        document = SimpleNamespace(iterate_items=lambda with_groups: iter([
            (title, 1), (noise, 1)]),
            pages={2: SimpleNamespace(size=SimpleNamespace(height=800))},
            export_to_dict=lambda: native)
        with patch("docling.document_converter.DocumentConverter") as converter:
            converter.return_value.convert.return_value = SimpleNamespace(
                status="success", document=document)
            items, _, warnings = parse_pdf(Path("unused.pdf"), {}, {"ocr": {"mode": "off"}})
        self.assertEqual(warnings, [])
        self.assertEqual([item["details"]["evidence_role"] for item in items],
                         ["authoritative", "structural"])

    def test_repeated_top_strip_is_running_header(self):
        header = "Position - recommandation et lignes directrices"
        pages = {str(page): {"size": {"height": 800}} for page in range(1, 4)}
        texts = [{"text": header, "prov": [{"page_no": page,
                  "bbox": {"l": 10, "t": 770, "r": 500, "b": 750,
                           "coord_origin": "BOTTOMLEFT"}, "charspan": [0, len(header)]}]}
                 for page in range(1, 4)]
        self.assertEqual(_pdf_running_headers({"pages": pages, "texts": texts}), {header})

    def test_running_header_mislabeled_as_list_item_is_structural(self):
        header = "Position - recommandation and publication details"
        pages = {page: SimpleNamespace(size=SimpleNamespace(width=600, height=800))
                 for page in range(1, 4)}
        nodes = [SimpleNamespace(label="list_item", text=header, marker=None,
                                 prov=[SimpleNamespace(page_no=page, bbox=SimpleNamespace(
                                     to_top_left_origin=lambda height: SimpleNamespace(
                                         l=20, t=30, r=500, b=50)))],
                                 self_ref=f"#/texts/{page}",
                                 parent=SimpleNamespace(cref="#/body"))
                 for page in pages]
        native = {"pages": {str(page): {"size": {"width": 600, "height": 800}}
                            for page in pages},
                  "texts": [{"self_ref": node.self_ref, "text": header,
                             "prov": [{"page_no": node.prov[0].page_no,
                                       "bbox": {"l": 20, "t": 750, "r": 500, "b": 730,
                                                "coord_origin": "BOTTOMLEFT"},
                                       "charspan": [0, len(header)]}]}
                            for node in nodes], "pictures": [], "tables": [], "groups": []}
        document = SimpleNamespace(iterate_items=lambda with_groups: iter((node, 1)
                                   for node in nodes), pages=pages, export_to_dict=lambda: native)
        with patch("docling.document_converter.DocumentConverter") as converter:
            converter.return_value.convert.return_value = SimpleNamespace(
                status="success", document=document)
            items, _, warnings = parse_pdf(Path("unused.pdf"), {}, {"ocr": {"mode": "off"}})
        self.assertEqual(warnings, [])
        self.assertEqual({item["kind"] for item in items}, {"paragraph"})
        self.assertEqual({item["details"]["evidence_role"] for item in items}, {"structural"})

    def test_running_header_inside_multipage_list_does_not_reclassify_body_fragments(self):
        header = 'Repeated publication header and edition information'
        text = 'c. The first clause continues. ' + header + ' d. The second clause continues.'
        def part(page, top, bottom, start, end):
            return SimpleNamespace(page_no=page, charspan=[start, end], bbox=SimpleNamespace(
                to_top_left_origin=lambda height: SimpleNamespace(l=40, t=top, r=500, b=bottom)))
        begin = text.index(header)
        node = SimpleNamespace(label='text', text=text, orig=text, self_ref='#/texts/0',
            parent=SimpleNamespace(cref='#/body'), prov=[part(1, 600, 620, 0, begin),
            part(2, 30, 50, begin, begin + len(header)),
            part(2, 80, 100, begin + len(header), len(text))])
        pages = {page: SimpleNamespace(size=SimpleNamespace(height=800, width=600)) for page in [1, 2, 3]}
        native = {'pages': {str(page): {'size': {'height': 800, 'width': 600}} for page in pages},
            'texts': [{'self_ref': '#/texts/0', 'text': text, 'prov': [
                {'page_no': p.page_no, 'charspan': p.charspan, 'bbox': {
                    'l': 40, 't': p.bbox.to_top_left_origin(800).t, 'r': 500,
                    'b': p.bbox.to_top_left_origin(800).b, 'coord_origin': 'TOPLEFT'}} for p in node.prov]}] +
            [{'text': header, 'prov': [{'page_no': page, 'bbox': {
                'l': 40, 't': 30, 'r': 500, 'b': 50, 'coord_origin': 'TOPLEFT'}}]} for page in [1, 3]],
            'pictures': [], 'tables': [], 'groups': []}
        # The repeated header witnesses need no emitted node for this adapter unit test.
        for index, witness in enumerate(native['texts'][1:], 1):
            witness['self_ref'] = f'#/texts/{index}'
            witness['content_layer'] = 'furniture'
        document = SimpleNamespace(pages=pages, export_to_dict=lambda: native,
                                   iterate_items=lambda with_groups: iter([(node, 1)]))
        converter = Mock()
        converter.convert.return_value = SimpleNamespace(status='success', document=document)
        items, _, warnings = parse_pdf(Path('unused.pdf'), {}, {'ocr': {'mode': 'off'}}, converter)
        self.assertEqual(warnings, [])
        self.assertEqual([(i['kind'], i['details']['evidence_role']) for i in items],
                         [('list_item', 'authoritative'), ('paragraph', 'structural'),
                          ('list_item', 'authoritative')])

    def test_short_running_headers_require_repeated_native_furniture_evidence(self):
        pages = {str(p): {'size': {'height': 800}} for p in [1, 2, 3]}
        texts = [{'text': 'XY', 'prov': [{'page_no': p, 'bbox': {
            'l': 40, 't': 30, 'r': 60, 'b': 50, 'coord_origin': 'TOPLEFT'}}]} for p in [1, 2, 3]]
        self.assertEqual(_pdf_running_headers({'pages': pages, 'texts': texts}), set())
        texts[0].update(label='page_header', content_layer='furniture')
        self.assertEqual(_pdf_running_headers({'pages': pages, 'texts': texts}), {'XY'})

    def test_native_heading_below_diagram_governs_an_aligned_table(self):
        def provenance(top, bottom):
            return [SimpleNamespace(page_no=1, bbox=SimpleNamespace(to_top_left_origin=lambda height:
                    SimpleNamespace(l=50, t=top, r=500, b=bottom)))]
        heading = SimpleNamespace(label='section_header', text='Investigations sent to the authority',
            self_ref='#/texts/0', parent=SimpleNamespace(cref='#/body'), prov=provenance(490, 510))
        cell = SimpleNamespace(text='Year', start_row_offset_idx=0, start_col_offset_idx=0,
            row_span=1, col_span=1, column_header=True, row_header=False, bbox=None)
        table = SimpleNamespace(label='table', self_ref='#/tables/0',
            parent=SimpleNamespace(cref='#/body'), prov=provenance(515, 600),
            data=SimpleNamespace(table_cells=[cell]))
        native = {'pages': {'1': {'size': {'height': 800, 'width': 600}}}, 'texts': [],
            'tables': [], 'groups': [], 'pictures': [{'self_ref': '#/pictures/0', 'children': [],
            'prov': [{'page_no': 1, 'bbox': {'l': 50, 't': 100, 'r': 500, 'b': 480,
                                           'coord_origin': 'TOPLEFT'}}]}]}
        document = SimpleNamespace(pages={1: SimpleNamespace(size=SimpleNamespace(height=800, width=600))},
            export_to_dict=lambda: native, iterate_items=lambda with_groups: iter([(heading, 1), (table, 1)]))
        converter = Mock()
        converter.convert.return_value = SimpleNamespace(status='success', document=document)
        items, _, warnings = parse_pdf(Path('unused.pdf'), {}, {'ocr': {'mode': 'off'}}, converter)
        self.assertEqual(warnings, [])
        self.assertEqual(items[0]['kind'], 'heading')
        self.assertEqual(items[0]['details']['evidence_role'], 'authoritative')
        self.assertEqual(items[1]['details']['heading_path'], [heading.text])

    def test_pdf_visual_caption_below_large_diagram_is_auxiliary(self):
        figure_box = {"l": 50, "t": 700, "r": 500, "b": 320,
                      "coord_origin": "BOTTOMLEFT"}
        label_box = {"l": 45, "t": 280, "r": 210, "b": 260,
                     "coord_origin": "BOTTOMLEFT"}
        picture = SimpleNamespace(label="picture", text="", marker=None,
                                  prov=[SimpleNamespace(page_no=23, bbox=SimpleNamespace(
                                      to_top_left_origin=lambda height: SimpleNamespace(
                                          l=50, t=100, r=500, b=480)))],
                                  self_ref="#/pictures/0", parent=SimpleNamespace(cref="#/body"))
        paragraph = SimpleNamespace(label="text", text="Infraction sous-jacente soupçonnée",
                                    marker=None, prov=[SimpleNamespace(page_no=23,
                                    bbox=SimpleNamespace(to_top_left_origin=lambda height:
                                    SimpleNamespace(l=45, t=520, r=210, b=540)))],
                                    self_ref="#/texts/0", parent=SimpleNamespace(cref="#/body"))
        native = {"pages": {"23": {"size": {"width": 600, "height": 800}}},
                  "pictures": [{"self_ref": "#/pictures/0", "prov": [
                      {"page_no": 23, "bbox": figure_box}], "children": []}],
                  "texts": [{"self_ref": "#/texts/0", "text": paragraph.text,
                             "parent": {"$ref": "#/body"}, "prov": [
                                 {"page_no": 23, "bbox": label_box}]}],
                  "tables": [], "groups": []}
        document = SimpleNamespace(iterate_items=lambda with_groups: iter([
            (picture, 1), (paragraph, 1)]),
            pages={23: SimpleNamespace(size=SimpleNamespace(height=800))},
            export_to_dict=lambda: native)
        with patch("docling.document_converter.DocumentConverter") as converter:
            converter.return_value.convert.return_value = SimpleNamespace(
                status="success", document=document)
            items, _, warnings = parse_pdf(Path("unused.pdf"), {}, {"ocr": {"mode": "off"}})
        self.assertEqual(warnings, [])
        self.assertEqual(items[1]["details"]["evidence_role"], "auxiliary")
        self.assertEqual(items[1]["parent_key"], items[0]["locator_key"])

    def test_short_heading_below_diagram_is_visual_annotation(self):
        figure_box = {"l": 50, "t": 700, "r": 500, "b": 320,
                      "coord_origin": "BOTTOMLEFT"}
        label_box = {"l": 45, "t": 280, "r": 210, "b": 260,
                     "coord_origin": "BOTTOMLEFT"}
        picture = SimpleNamespace(label="picture", text="", marker=None,
                                  prov=[SimpleNamespace(page_no=2, bbox=SimpleNamespace(
                                      to_top_left_origin=lambda height: SimpleNamespace(
                                          l=50, t=100, r=500, b=480)))],
                                  self_ref="#/pictures/0", parent=SimpleNamespace(cref="#/body"))
        heading = SimpleNamespace(label="section_header", text="Diagram labels", marker=None,
                                  prov=[SimpleNamespace(page_no=2, bbox=SimpleNamespace(
                                      to_top_left_origin=lambda height: SimpleNamespace(
                                          l=45, t=520, r=210, b=540)))],
                                  self_ref="#/texts/0", parent=SimpleNamespace(cref="#/body"))
        native = {"pages": {"2": {"size": {"width": 600, "height": 800}}},
                  "pictures": [{"self_ref": "#/pictures/0", "prov": [{"page_no": 2,
                      "bbox": figure_box}], "children": []}],
                  "texts": [{"self_ref": "#/texts/0", "text": heading.text,
                             "parent": {"$ref": "#/body"}, "prov": [{"page_no": 2,
                                 "bbox": label_box}]}], "tables": [], "groups": []}
        document = SimpleNamespace(iterate_items=lambda with_groups: iter([
            (picture, 1), (heading, 1)]),
            pages={2: SimpleNamespace(size=SimpleNamespace(height=800))},
            export_to_dict=lambda: native)
        with patch("docling.document_converter.DocumentConverter") as converter:
            converter.return_value.convert.return_value = SimpleNamespace(
                status="success", document=document)
            items, _, warnings = parse_pdf(Path("unused.pdf"), {}, {"ocr": {"mode": "off"}})
        self.assertEqual(warnings, [])
        self.assertEqual([item["kind"] for item in items], ["figure", "paragraph"])
        self.assertEqual(items[1]["details"]["evidence_role"], "auxiliary")

    def test_pdf_numbered_quotation_lead_and_lettered_alternative(self):
        def node(label, value, ref):
            return SimpleNamespace(label=label, text=value, marker=None,
                                   prov=[SimpleNamespace(page_no=2, bbox=None)],
                                   self_ref=ref, parent=SimpleNamespace(cref="#/body"))
        lead = node("section_header", "219. Le client est:", "#/texts/0")
        option = node("text", "a. une personne physique ou morale", "#/texts/1")
        document = SimpleNamespace(iterate_items=lambda with_groups: iter([
            (lead, 1), (option, 1)]),
            pages={2: SimpleNamespace(size=SimpleNamespace(height=800))},
            export_to_dict=lambda: {"texts": [], "pictures": [], "tables": [], "groups": []})
        with patch("docling.document_converter.DocumentConverter") as converter:
            converter.return_value.convert.return_value = SimpleNamespace(
                status="success", document=document)
            items, _, _ = parse_pdf(Path("unused.pdf"), {}, {"ocr": {"mode": "off"}})
        self.assertEqual([item["kind"] for item in items], ["paragraph", "list_item"])

    def test_sentence_like_ellipsis_heading_becomes_paragraph(self):
        node = SimpleNamespace(label="section_header", text="Par exception, les diligences sont conduites…",
                               marker=None, prov=[SimpleNamespace(page_no=1, bbox=None)],
                               self_ref="#/texts/0", parent=SimpleNamespace(cref="#/body"))
        document = SimpleNamespace(iterate_items=lambda with_groups: iter([(node, 1)]),
                                   pages={1: SimpleNamespace(size=SimpleNamespace(height=800))},
                                   export_to_dict=lambda: {"texts": [], "pictures": [],
                                                           "tables": [], "groups": []})
        with patch("docling.document_converter.DocumentConverter") as converter:
            converter.return_value.convert.return_value = SimpleNamespace(
                status="success", document=document)
            items, _, _ = parse_pdf(Path("unused.pdf"), {}, {"ocr": {"mode": "off"}})
        self.assertEqual(items[0]["kind"], "paragraph")

    def test_pdf_figure_asset_is_deterministic_and_linked(self):
        from PIL import Image, ImageDraw

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            image = Image.new("RGB", (120, 120), "white")
            ImageDraw.Draw(image).rectangle((20, 20, 80, 80), outline="black", width=3)
            pdf = root / "source.pdf"
            image.save(pdf, "PDF", resolution=72)
            element = {"id": "el:" + "a" * 64, "kind": "figure",
                       "location": {"format": "pdf", "page": 1,
                                    "bbox": [20, 20, 80, 80]}}
            outputs = []
            for index in (1, 2):
                target = root / f"run-{index}"
                target.mkdir()
                sidecar = [{"element_id": element["id"], "docling_ref": "#/pictures/0"}]
                warnings = []
                _save_pdf_figure_assets(pdf, [element], sidecar, target, warnings)
                asset = sidecar[0]["asset"]
                data = (target / asset["path"]).read_bytes()
                self.assertEqual(warnings, [])
                self.assertEqual(asset["media_type"], "image/png")
                self.assertEqual(asset["representation_role"], "authoritative_visual")
                self.assertEqual(asset["sha256"], hashlib.sha256(data).hexdigest())
                self.assertEqual(Image.open(io.BytesIO(data)).size, (120, 120))
                outputs.append((asset, data))
            self.assertEqual(outputs[0], outputs[1])

    def test_pdf_parse_output_links_figure_asset_without_schema_change(self):
        from PIL import Image

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "diagram.pdf"
            Image.new("RGB", (120, 120), "white").save(path, "PDF", resolution=72)
            entry = _fixture(root, "diagram.pdf", "application/pdf", path.read_bytes())
            config = _config(root, [entry])
            figure = {"kind": "figure", "text": "", "locator_key": "pdf:p1:block-1",
                      "location": {"format": "pdf", "page": 1, "label": "block 1",
                                   "bbox": [20, 20, 80, 80]},
                      "parent_key": None,
                      "details": {"docling_ref": "#/pictures/0", "evidence_role": "structural"}}
            native = {"pages": {"1": {"size": {"width": 120, "height": 120}}},
                      "pictures": [{"self_ref": "#/pictures/0"}], "texts": [], "tables": []}
            with patch.dict(PARSERS, {"pdf": (lambda *_args: ([figure], native, []), "fixture")}):
                outcome = parse_collection(config, root / "output")["outcomes"][0]
            self.assertEqual(outcome["status"], "parsed")
            folder = Path(outcome["output"])
            bundle = json.loads((folder / "artifact-bundle.json").read_text())
            sidecar = json.loads((folder / "parse.json").read_text())
            Draft202012Validator(SCHEMA, format_checker=FormatChecker()).validate(bundle)
            self.assertEqual(sidecar["elements"][0]["element_id"], bundle["elements"][0]["id"])
            asset = sidecar["elements"][0]["asset"]
            self.assertEqual(asset["sha256"], hashlib.sha256((folder / asset["path"]).read_bytes()).hexdigest())

    def _picture_fixture(self, page, nodes):
        picture_ref = "#/pictures/0"
        native_nodes = []
        for index, entry in enumerate(nodes):
            label, value, box = entry[:3]
            native_nodes.append({"self_ref": f"#/texts/{index}", "label": label,
                                 "text": value, "orig": entry[3] if len(entry) > 3 else value,
                                 "content_layer": "body",
                                 "parent": {"$ref": picture_ref},
                                 "prov": [{"page_no": page, "bbox": dict(zip(
                                     ("l", "t", "r", "b"), box))}]})
            if len(entry) > 4:
                native_nodes[-1]['prov'][0]['bbox']['coord_origin'] = entry[4]
        native = {"pages": {str(page): {"size": {"width": 600, "height": 842}}},
                  "texts": native_nodes, "groups": [], "tables": [], "pictures": [
            {"self_ref": picture_ref, "children": [
                {"$ref": node["self_ref"]} for node in native_nodes]}]}
        picture = SimpleNamespace(label="picture", text="", marker=None,
                                  prov=[SimpleNamespace(page_no=page, bbox=None)],
                                  self_ref=picture_ref, parent=SimpleNamespace(cref="#/body"))
        document = SimpleNamespace(iterate_items=lambda with_groups: iter([(picture, 1)]),
                                   pages={page: SimpleNamespace(size=SimpleNamespace(height=842))},
                                   export_to_dict=lambda: native)
        with patch("docling.document_converter.DocumentConverter") as converter:
            converter.return_value.convert.return_value = SimpleNamespace(
                status="success", document=document)
            items, _, warnings = parse_pdf(Path("unused.pdf"), {}, {"ocr": {"mode": "off"}})
        return items, warnings, native

    def test_picture_fragments_with_mixed_coordinate_origins_keep_text_and_source_parts(self):
        items, warnings, native = self._picture_fixture(1, [
            ('text', 'Sensor readings from', [40, 100, 150, 110], 'Sensor readings from', 'TOPLEFT'),
            ('text', 'the calibration device.', [155, 742, 280, 732], 'the calibration device.', 'BOTTOMLEFT')])
        combined = next(item for item in items if item['text'] == 'Sensor readings from the calibration device.')
        self.assertEqual(combined['location']['bbox'], [40, 100, 280, 110])
        self.assertEqual(combined['details']['source_parts'][1]['provenance'], native['texts'][1]['prov'])
        self.assertFalse(warnings)

    def test_short_picture_date_heading_is_retained_without_an_ocr_warning(self):
        items, warnings, _ = self._picture_fixture(1, [
            ('section_header', 'January 2031', [40, 740, 200, 720])])
        self.assertTrue(any(item['kind'] == 'heading' and item['text'] == 'January 2031' for item in items))
        self.assertFalse(warnings)

    def test_pdf_picture_short_list_items_a_page_21(self):
        items, warnings, _ = self._picture_fixture(21, [
            ("text", "L'organisme doit pouvoir identifier les retraits effectués dans des conditions particulièrement inhabituelles et documenter les opérations.",
             (72, 758, 531, 705)),
            ("text", "Les informations recueillies au moment de l'opération permettent de circonstancier une déclaration transmise à Tracfin.",
             (72, 686, 531, 647)),
            ("list_item", "date et heure de l'opération ;", (90, 636, 249, 626),
             "- date et heure de l'opération ;"),
            ("text", "- lieux de l'opération ;", (90, 622, 207, 612)),
            ("list_item", "référence de la carte utilisée", (90, 607, 247, 597),
             "- référence de la carte utilisée"),
            ("text", "- montant et composition de la somme requise.", (90, 592, 345, 582)),
        ])
        self.assertEqual([item["text"] for item in items[-4:]],
                         ["- date et heure de l'opération ;", "- lieux de l'opération ;",
                          "- référence de la carte utilisée",
                          "- montant et composition de la somme requise."])
        self.assertEqual([item["kind"] for item in items[-4:]], ["list_item"] * 4)
        self.assertEqual([item["details"]["evidence_role"] for item in items[1:]],
                         ["authoritative"] * 6)
        self.assertEqual([item["parent_key"] for item in items[1:]], [None] * 6)
        self.assertEqual(warnings, [])

    def test_pdf_picture_word_fragments_a_page_63(self):
        words = ["Je", "joins", "obligatoirement", "le", "formulaire", "de", "désignation."]
        boxes = [(197, 416, 205, 404), (209, 416, 227, 404), (229, 415, 413, 405),
                 (296, 416, 304, 404), (308, 416, 349, 404), (353, 416, 362, 404),
                 (366, 416, 415, 404)]
        branches = [("text", "Le représentant légal est une personne physique", (92, 276, 190, 258)),
                    ("text", "Le représentant légal est une personne morale", (214, 276, 330, 258))]
        items, warnings, native = self._picture_fixture(
            63, [("text", word, box) for word, box in zip(words, boxes)] + branches)
        self.assertEqual(items[1]["text"], "Je joins obligatoirement le formulaire de désignation.")
        self.assertEqual(items[1]["details"]["docling_refs"],
                         [node["self_ref"] for node in native["texts"][:7]])
        self.assertEqual([part["text"] for part in items[1]["details"]["source_parts"]], words)
        self.assertEqual(items[1]["location"]["page"], 63)
        self.assertEqual(items[1]["location"]["bbox"], [197, 426, 415, 438])
        self.assertEqual(items[1]["details"]["evidence_role"], "auxiliary")
        self.assertEqual([item["text"] for item in items[2:]], [entry[1] for entry in branches])
        self.assertEqual(warnings, [])
        repeat_items, repeat_warnings, _ = self._picture_fixture(
            63, [("text", word, box) for word, box in zip(words, boxes)] + branches)
        self.assertEqual((repeat_items, repeat_warnings), (items, warnings))

    def test_mixed_picture_keeps_document_flow_before_diagram(self):
        items, warnings, _ = self._picture_fixture(63, [
            ("text", "Un déclarant est une personne habilitée à établir et transmettre une déclaration à TRACFIN pour son organisme.",
             (88, 557, 516, 519)),
            ("text", "Un correspondant est une personne habilitée à renseigner TRACFIN en cas de demande d'information sur un dossier.",
             (88, 505, 516, 476)),
            ("section_header", "Schéma récapitulatif", (89, 464, 187, 453)),
            ("text", "Je joins obligatoirement le formulaire de désignation.",
             (196, 416, 415, 404)),
        ])
        self.assertEqual(warnings, [])
        self.assertEqual([item["details"]["evidence_role"] for item in items[1:]],
                         ["authoritative", "authoritative", "authoritative", "auxiliary"])
        self.assertEqual(items[-1]["parent_key"], items[0]["locator_key"])

    def test_pdf_picture_diagram_phrases_t_page_37_and_graphical_noise(self):
        items, warnings, _ = self._picture_fixture(37, [
            ("text", "Achat d'un bien", (282, 214, 344, 207)),
            ("text", "immobilier pour 1,5 M€", (267, 205, 356, 198)),
            ("text", "Société A", (202, 244, 239, 237)),
            ("text", "Acquisition", (114, 214, 158, 207)),
            ("text", "de titres pour 3 M€", (98, 205, 172, 198)),
            ("text", "|", (144, 434, 145, 419)),
            ("text", "6 M€", (210, 378, 229, 371)),
        ])
        self.assertEqual([item["text"] for item in items[1:]],
                         ["Achat d'un bien immobilier pour 1,5 M€",
                          "Acquisition de titres pour 3 M€"])
        self.assertEqual([item["details"]["evidence_role"] for item in items[1:]],
                         ["auxiliary", "auxiliary"])
        self.assertEqual(warnings[0]["code"], "excluded_picture_ocr_fragments")
        self.assertEqual(warnings[0]["refs"], ["#/texts/2", "#/texts/5", "#/texts/6"])

    def test_long_equipment_diagram_labels_do_not_establish_document_prose(self):
        note = ("The equipment labels in this schematic describe routing only and do not establish "
                "which maintenance team is authorised to approve a replacement request.")
        items, warnings, _ = self._picture_fixture(1, [
            ("footnote", note, (20, 100, 550, 70)),
            ("text", "Department responsible for calibration and inspection of all portable instruments used by the northern observatory", (20, 700, 180, 660)),
            ("text", "Committee responsible for the allocation and replacement of all laboratory equipment at the southern observatory", (300, 700, 460, 660)),
            ("text", "Operators responsible for the inspection and maintenance of", (20, 600, 180, 580)),
        ])
        recovered_note = next(item for item in items if item["text"] == note)
        self.assertEqual(recovered_note["details"]["evidence_role"], "authoritative")
        labels = [item for item in items if item["kind"] != "figure" and item["text"] != note]
        self.assertTrue(labels)
        self.assertTrue(all(item["details"]["evidence_role"] == "auxiliary" for item in labels))
        self.assertNotIn("potential_picture_text", {warning["code"] for warning in warnings})
        self.assertIn("excluded_figure_text_uncertain", {warning["code"] for warning in warnings})

    def test_unfinished_picture_prose_stays_pending_amid_real_document_flow(self):
        items, warnings, _ = self._picture_fixture(1, [
            ("text", "The maintenance department records each instrument inspection and retains the results until the instrument is retired.", (20, 700, 550, 670)),
            ("text", "The observatory director reviews these records every month and approves any necessary changes to the inspection schedule.", (20, 640, 550, 610)),
            ("text", "The audit operator must retain all maintenance records for the lifetime of", (20, 580, 550, 560)),
        ])
        self.assertEqual(sum(item["details"]["evidence_role"] == "authoritative" for item in items), 2)
        self.assertIn("potential_picture_text", {warning["code"] for warning in warnings})

    def test_pdf_picture_two_word_label_and_unresolved_fragment(self):
        items, warnings, _ = self._picture_fixture(1, [
            ("caption", "Compte bancaire", (10, 100, 100, 90)),
            ("text", "Actionnaire à 100 %", (10, 90, 100, 80)),
            ("list_item", "LeKBIS demon de", (10, 80, 100, 70)),
            ("text", "O", (10, 60, 20, 50)),
        ])
        self.assertEqual([item["text"] for item in items[1:]],
                         ["Compte bancaire", "Actionnaire à 100 %"])
        self.assertEqual({warning["code"] for warning in warnings},
                         {"excluded_figure_text_uncertain", "excluded_picture_ocr_fragments"})

    def test_corrupt_picture_ocr_and_publisher_footer_are_excluded(self):
        corrupt = "“FT voddns à] 2219281007 é 23/qOuId Un ZaQUCoUa SOA"
        items, warnings, _ = self._picture_fixture(63, [
            ("text", corrupt, (388, 98, 516, 80)),
        ])
        self.assertEqual([item["kind"] for item in items], ["figure"])
        self.assertEqual([warning["code"] for warning in warnings],
                         ["excluded_picture_corrupt_ocr"])
        items, warnings, _ = self._picture_fixture(82, [
            ("text", "Suivez Tracfin sur", (273, 96, 360, 85)),
            ("text", "www.economie.gouv.fr/tracfin www.linkedin.com/company/tracfin",
             (274, 84, 448, 60)),
        ])
        self.assertEqual([item["kind"] for item in items], ["figure"])
        self.assertEqual([warning["code"] for warning in warnings],
                         ["excluded_picture_furniture"])
        self.assertEqual(warnings[0]["refs"], ["#/texts/0", "#/texts/1"])

    def test_pdf_document_index_keeps_contents_and_cells(self):
        cell = SimpleNamespace(start_row_offset_idx=0, start_col_offset_idx=0,
                               text="INTRODUCTION", bbox=None, row_span=1, col_span=1,
                               column_header=False, row_header=False)
        index = SimpleNamespace(label="document_index", text="", marker=None,
                                prov=[SimpleNamespace(page_no=2, bbox=None)],
                                data=SimpleNamespace(table_cells=[cell]),
                                self_ref="#/tables/0", parent=SimpleNamespace(cref="#/body"),
                                export_to_markdown=lambda doc: "| INTRODUCTION | 2 |")
        document = SimpleNamespace(
            iterate_items=lambda with_groups: iter([(index, 1)]),
            pages={2: SimpleNamespace(size=SimpleNamespace(height=842))},
            export_to_dict=lambda: {},
        )
        with patch("docling.document_converter.DocumentConverter") as converter:
            converter.return_value.convert.return_value = SimpleNamespace(status="success", document=document)
            items, _, warnings = parse_pdf(Path("unused.pdf"), {}, {"ocr": {"mode": "off"}})
        self.assertEqual(warnings, [])
        self.assertEqual([item["kind"] for item in items], ["table", "table_cell"])
        self.assertIn("INTRODUCTION", items[0]["text"])
        self.assertEqual(items[1]["text"], "INTRODUCTION")
        self.assertEqual(items[1]["parent_key"], items[0]["locator_key"])
        self.assertEqual(items[0]["location"]["page"], 2)

    def test_pdf_picture_children_charspans_and_heading_hierarchy(self):
        def part(page, start, end):
            return SimpleNamespace(page_no=page, bbox=None, charspan=[start, end])

        picture = SimpleNamespace(label="picture", text="", marker=None,
                                  prov=[part(1, 0, 0)], self_ref="#/pictures/0",
                                  parent=SimpleNamespace(cref="#/body"))
        heading = SimpleNamespace(label="section_header", text="I. Main section", marker=None,
                                  prov=[part(1, 0, 15)], self_ref="#/texts/0",
                                  parent=SimpleNamespace(cref="#/body"))
        subsection = SimpleNamespace(label="section_header", text="1.1. Details", marker=None,
                                     prov=[part(1, 0, 12)], self_ref="#/texts/1",
                                     parent=SimpleNamespace(cref="#/body"))
        source_text = "First page has a meaningful passage. Second page has another meaningful passage."
        split = source_text.index("Second")
        paragraph = SimpleNamespace(label="text", text=source_text, marker=None,
                                    prov=[part(1, 0, split), part(2, split, len(source_text))],
                                    self_ref="#/texts/2", parent=SimpleNamespace(cref="#/body"))
        native = {"texts": [
            {"self_ref": "#/texts/3", "label": "text", "content_layer": "body",
             "text": "Meaningful text inside this pictured process flow remains available.",
             "prov": [{"page_no": 2, "charspan": [0, 66]}]},
            {"self_ref": "#/texts/4", "label": "text", "content_layer": "body",
             "text": "x", "prov": [{"page_no": 2, "charspan": [0, 1]}]},
        ], "pictures": [{"self_ref": "#/pictures/0", "children": [
            {"$ref": "#/groups/0"}, {"$ref": "#/texts/4"}]}],
            "groups": [{"self_ref": "#/groups/0", "children": [{"$ref": "#/texts/3"}]}],
            "tables": []}
        document = SimpleNamespace(
            iterate_items=lambda with_groups: iter([(heading, 1), (subsection, 1),
                                                     (paragraph, 1), (picture, 1)]),
            pages={1: SimpleNamespace(size=SimpleNamespace(height=842)),
                   2: SimpleNamespace(size=SimpleNamespace(height=842))},
            export_to_dict=lambda: native)
        with patch("docling.document_converter.DocumentConverter") as converter:
            converter.return_value.convert.return_value = SimpleNamespace(status="success", document=document)
            items, _, warnings = parse_pdf(Path("unused.pdf"), {}, {"ocr": {"mode": "off"}})
        self.assertEqual([item["text"] for item in items
                          if item["details"].get("docling_ref") == "#/texts/2"],
                         [source_text[:split].strip(), source_text[split:].strip()])
        self.assertEqual([item["location"]["page"] for item in items],
                         sorted(item["location"]["page"] for item in items))
        self.assertEqual(items[1]["details"]["level"], 2)
        self.assertEqual(items[1]["parent_key"], items[0]["locator_key"])
        picture_key = next(item["locator_key"] for item in items if item["kind"] == "figure")
        picture_prose = next(item for item in items
                             if item["details"].get("docling_ref") == "#/texts/3")
        self.assertEqual(picture_prose["parent_key"], picture_key)
        self.assertEqual(picture_prose["kind"], "paragraph")
        self.assertEqual([warning["code"] for warning in warnings], ["excluded_picture_ocr_fragments"])
        self.assertEqual(warnings[0]["refs"], ["#/texts/4"])

    def test_pdf_empty_table_is_diagnostic_and_checkbox_keeps_state(self):
        provenance = SimpleNamespace(page_no=1, bbox=None, charspan=[0, 22])
        empty_table = SimpleNamespace(label="table", text="", marker=None,
                                      prov=[provenance], data=SimpleNamespace(table_cells=[]),
                                      self_ref="#/tables/0", parent=SimpleNamespace(cref="#/body"),
                                      export_to_markdown=lambda doc: "")
        checkbox = SimpleNamespace(label="checkbox_unselected", text="Review the account owner", marker=None,
                                   prov=[provenance], self_ref="#/texts/0",
                                   parent=SimpleNamespace(cref="#/body"))
        document = SimpleNamespace(
            iterate_items=lambda with_groups: iter([(empty_table, 1), (checkbox, 1)]),
            pages={1: SimpleNamespace(size=SimpleNamespace(height=842))},
            export_to_dict=lambda: {"texts": [{"self_ref": "#/texts/0"}],
                                    "tables": [{"self_ref": "#/tables/0"}]})
        with patch("docling.document_converter.DocumentConverter") as converter:
            converter.return_value.convert.return_value = SimpleNamespace(status="success", document=document)
            items, _, warnings = parse_pdf(Path("unused.pdf"), {}, {"ocr": {"mode": "off"}})
        self.assertEqual([item["kind"] for item in items], ["list_item"])
        self.assertEqual(items[0]["details"]["checkbox_state"], "unselected")
        self.assertEqual([warning["code"] for warning in warnings], ["empty_docling_table"])

    def test_pilot_html_clauses_have_resolvable_selectors_and_metadata_origins(self):
        with tempfile.TemporaryDirectory() as directory:
            report = parse_collection(ROOT / "config/pipeline.json", directory, ["L1", "L2"])
            self.assertEqual([item["status"] for item in report["outcomes"]], ["parsed", "parsed"])
            for outcome in report["outcomes"]:
                folder = Path(outcome["output"])
                bundle = json.loads((folder / "artifact-bundle.json").read_text())
                sidecar = json.loads((folder / "parse.json").read_text())
                Draft202012Validator(SCHEMA, format_checker=FormatChecker()).validate(bundle)
                document = html.parse(str(ROOT / outcome["local_path"])).getroot()
                self.assertTrue(bundle["elements"])
                for element in bundle["elements"]:
                    matches = document.cssselect(element["location"]["selector"])
                    self.assertEqual(len(matches), 1)
                    self.assertIn(element["text"], " ".join(" ".join(matches[0].itertext()).split()))
                self.assertEqual(sidecar["metadata_origins"]["source.version"]["origin"], "manifest")
                self.assertEqual(sidecar["metadata_origins"]["source.id"]["origin"], "manifest")
                self.assertEqual(sidecar["metadata_origins"]["html.content_selector"]["origin"], "parsing_config")
                self.assertEqual(len(sidecar["elements"]), len(bundle["elements"]))
            l1 = json.loads((Path(report["outcomes"][0]["output"]) / "artifact-bundle.json").read_text())
            self.assertEqual({element["location"].get("clause") for element in l1["elements"]
                              if "clause" in element["location"]}, {"1°", "2°"})

    def test_html_source_override_preserves_article_location(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = json.loads((ROOT / "config/pipeline.json").read_text(encoding="utf-8"))
            config.update(root=str(ROOT), manifest_ids=["L1"])
            config["parsing"]["html"] = {"content_selector": "main", "source_overrides": {
                "L1": {"content_selector": "article.js-article-content-to-copy", "article_id": "test-article"}}}
            path = root / "config.json"
            path.write_text(json.dumps(config), encoding="utf-8")
            outcome = parse_collection(path, root / "output")["outcomes"][0]
            self.assertEqual(outcome["status"], "parsed")
            folder = Path(outcome["output"])
            bundle = json.loads((folder / "artifact-bundle.json").read_text())
            sidecar = json.loads((folder / "parse.json").read_text())
            self.assertTrue(all(e["location"]["article_id"] == "test-article" for e in bundle["elements"]))
            self.assertEqual(sidecar["metadata_origins"]["html.article_id"]["origin"], "parsing_config")

    def test_markdown_and_docx_preserve_order_tables_and_locations(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            markdown = "# Heading\n\nFirst paragraph.\n\n| A | B |\n|---|---|\n| x | y |\n"
            md = _fixture(root, "example.md", "text/markdown", markdown)
            document = Document()
            document.add_heading("Heading", level=1)
            document.add_paragraph("First paragraph.")
            table = document.add_table(rows=2, cols=2)
            for row, values in zip(table.rows, (("A", "B"), ("x", "y"))):
                for cell, value in zip(row.cells, values):
                    cell.text = value
            docx_path = root / "example.docx"
            document.save(docx_path)
            docx = _fixture(root, "example.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            docx_path.read_bytes())
            config = _config(root, [md, docx])
            report = parse_collection(config, root / "output")
            self.assertEqual([item["status"] for item in report["outcomes"]], ["parsed", "parsed"])
            for outcome in report["outcomes"]:
                folder = Path(outcome["output"])
                bundle = json.loads((folder / "artifact-bundle.json").read_text())
                sidecar = json.loads((folder / "parse.json").read_text())
                self.assertEqual([e["ordinal"] for e in bundle["elements"]], list(range(7)))
                self.assertEqual([e["kind"] for e in bundle["elements"][:3]],
                                 ["heading", "paragraph", "table"])
                self.assertEqual(len([e for e in bundle["elements"] if e["kind"] == "table_cell"]), 4)
                self.assertEqual(len([e for e in sidecar["elements"][3:] if e["parent_element_id"]]), 4)
                self.assertEqual(sidecar["elements"][1]["parent_element_id"], bundle["elements"][0]["id"])
                Draft202012Validator(SCHEMA, format_checker=FormatChecker()).validate(bundle)
            md_bundle = json.loads((Path(report["outcomes"][0]["output"]) / "artifact-bundle.json").read_text())
            self.assertEqual(md_bundle["elements"][1]["location"]["line_start"], 3)
            docx_bundle = json.loads((Path(report["outcomes"][1]["output"]) / "artifact-bundle.json").read_text())
            self.assertEqual(docx_bundle["elements"][2]["location"]["block_index"], 2)

    def test_bad_source_does_not_block_good_source(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            good = _fixture(root, "good.md", "text/markdown", "# Heading\n\nText.\n")
            missing = dict(good, manifest_id="missing.md", local_path="missing.md")
            config = _config(root, [missing, good])
            report = parse_collection(config, root / "output")
            self.assertEqual([item["status"] for item in report["outcomes"]], ["failed", "parsed"])
            self.assertEqual(report["outcomes"][0]["error"]["code"], "missing_file")
            self.assertTrue((root / "output" / "report.json").is_file())

    def test_cli_source_must_be_in_configured_selection(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            first = _fixture(root, "first.md", "text/markdown", "# First\n")
            second = _fixture(root, "second.md", "text/markdown", "# Second\n")
            config_path = _config(root, [first, second])
            config = json.loads(config_path.read_text())
            config["manifest_ids"] = ["first.md"]
            config_path.write_text(json.dumps(config))
            with self.assertRaisesRegex(ValueError, "not configured"):
                parse_collection(config_path, root / "output", ["second.md"])

    def test_repeated_runs_replace_report(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            entry = _fixture(root, "good.md", "text/markdown", "# Heading\n")
            config = _config(root, [entry])
            output = root / "output"
            first = parse_collection(config, output)
            second = parse_collection(config, output)

            self.assertNotEqual(first["run_id"], second["run_id"])
            self.assertEqual(len(list(output.glob("report-*.json"))), 0)
            for report in (first, second):
                self.assertRegex(report["run_id"], r"^\d{8}-\d{2}-\d{2}-\d{2}-[0-9a-f]{8}$")
                self.assertIn("+00:00", report["started_at"])
            self.assertEqual(json.loads((output / "report.json").read_text()), second)

    def test_cli_reports_each_document_and_continues_after_parser_error(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            bad = _fixture(root, "bad.md", "text/markdown", "# Broken\n")
            good = _fixture(root, "good.md", "text/markdown", "# Heading\n\nText.\n")
            config = _config(root, [bad, good])
            stdout, stderr = io.StringIO(), io.StringIO()

            def parse_or_fail(path, source, settings):
                if path.name == "bad.md":
                    raise ValueError("broken document")
                return parse_markdown(path, source, settings)

            with patch.dict(PARSERS, {"markdown": (parse_or_fail, "fixture")}), \
                    redirect_stdout(stdout), redirect_stderr(stderr):
                exit_code = main(["--config", str(config), "--output-dir", str(root / "output")])

            self.assertEqual(exit_code, 1)
            report = json.loads(stdout.getvalue())
            self.assertEqual([item["status"] for item in report["outcomes"]], ["failed", "parsed"])
            self.assertEqual([item["status"] for item in json.loads(
                (root / "output" / "report.json").read_text())["outcomes"]], ["failed", "parsed"])
            self.assertEqual(len(stderr.getvalue().splitlines()), 5)
            progress = stderr.getvalue().splitlines()
            self.assertIn("[1/2] bad.md bad.md: parsing...", progress[0])
            self.assertIn("[1/2] bad.md bad.md: failed in ", progress[1])
            self.assertIn("broken document", progress[1])
            self.assertIn("[2/2] good.md good.md: parsing...", progress[2])
            self.assertIn("[2/2] good.md good.md: succeeded in ", progress[3])
            self.assertIn("(2 elements)", progress[3])
            self.assertIn("report.json", progress[4])
            self.assertGreaterEqual(report["outcomes"][1]["elapsed_seconds"], 0)

    def test_docx_nested_table_keeps_parent_and_unique_location(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "nested.docx"
            document = Document()
            outer = document.add_table(rows=1, cols=1)
            nested = outer.cell(0, 0).add_table(rows=1, cols=1)
            nested.cell(0, 0).text = "Nested evidence"
            document.save(path)
            items, _native, _warnings = parse_docx(path, {}, {})
            self.assertEqual([item["kind"] for item in items],
                             ["table", "table_cell", "table", "table_cell"])
            self.assertEqual(items[2]["parent_key"], items[1]["locator_key"])
            self.assertEqual(items[3]["location"]["table_index"], 1)

    def test_docx_vertical_merge_keeps_one_cell_with_row_span(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "merged.docx"
            document = Document()
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Merged"
            table.cell(0, 1).text = "A"
            table.cell(1, 1).text = "B"
            table.cell(0, 0).merge(table.cell(1, 0))
            document.save(path)
            items, _native, _warnings = parse_docx(path, {}, {})
            cells = [item for item in items if item["kind"] == "table_cell"]
            self.assertEqual([item["text"] for item in cells], ["Merged", "A", "B"])
            self.assertEqual(cells[0]["details"]["row_span"], 2)

    def test_missing_tesseract_fails_pdf_and_continues_to_next_document(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pdf = _fixture(root, "scan.pdf", "application/pdf", b"%PDF-1.7\nfixture")
            markdown = _fixture(root, "good.md", "text/markdown", "# Heading\n")
            config_path = _config(root, [pdf, markdown])
            config = json.loads(config_path.read_text())
            config["parsing"]["ocr"] = {"mode": "auto", "tesseract_cmd": "kf-no-such-tesseract"}
            config_path.write_text(json.dumps(config))
            report = parse_collection(config_path, root / "output")
            self.assertEqual([item["status"] for item in report["outcomes"]], ["failed", "parsed"])
            self.assertEqual(report["outcomes"][0]["error"]["code"], "ocr_setup_error")
            self.assertIn("Tesseract CLI unavailable", report["outcomes"][0]["error"]["message"])

    def test_partial_parse_retains_valid_elements_and_warning(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            entry = _fixture(root, "partial.md", "text/markdown", "# Located\n")
            config = _config(root, [entry])
            item = {"kind": "heading", "text": "Located", "locator_key": "markdown:line-1:block-0",
                    "location": {"format": "markdown", "heading_path": ["Located"],
                                 "line_start": 1, "line_end": 1},
                    "parent_key": None, "details": {}}
            with patch.dict(PARSERS, {"markdown": (lambda *_args: (
                [item], None, [{"code": "unlocated_block", "message": "omitted content"}]), "fixture")}):
                report = parse_collection(config, root / "output")
            outcome = report["outcomes"][0]
            self.assertEqual(outcome["status"], "partial")
            self.assertEqual(outcome["warnings"][0]["code"], "unlocated_block")
            bundle = json.loads((Path(outcome["output"]) / "artifact-bundle.json").read_text())
            self.assertEqual(len(bundle["elements"]), 1)

    def test_documented_ocr_exclusion_does_not_make_parse_partial(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            entry = _fixture(root, "sample.md", "text/markdown", "# Located\n")
            config = _config(root, [entry])
            item = {"kind": "heading", "text": "Located", "locator_key": "markdown:line-1:block-0",
                    "location": {"format": "markdown", "heading_path": ["Located"],
                                 "line_start": 1, "line_end": 1},
                    "parent_key": None, "details": {}}
            warning = {"code": "excluded_picture_ocr_fragments", "refs": ["#/texts/0"]}
            with patch.dict(PARSERS, {"markdown": (lambda *_args: ([item], None, [warning]), "fixture")}):
                outcome = parse_collection(config, root / "output")["outcomes"][0]
            self.assertEqual(outcome["status"], "parsed")
            self.assertEqual(outcome["warnings"], [warning])

    def test_retained_ocr_review_requirement_does_not_make_parse_partial(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            entry = _fixture(root, 'sample.md', 'text/markdown', '# Located\n')
            config = _config(root, [entry])
            item = {'kind': 'heading', 'text': 'Located', 'locator_key': 'markdown:line-1:block-0',
                    'location': {'format': 'markdown', 'heading_path': ['Located'],
                                 'line_start': 1, 'line_end': 1}, 'parent_key': None, 'details': {}}
            warning = {'code': 'source_ocr_verification_uncertain', 'ref': '#/texts/0', 'page': 1}
            with patch.dict(PARSERS, {'markdown': (lambda *_: ([item], None, [warning]), 'fixture')}):
                outcome = parse_collection(config, root / 'output')['outcomes'][0]
            self.assertEqual(outcome['status'], 'parsed')
            self.assertEqual(outcome['warnings'], [warning])

    def test_markdown_code_block_keeps_text_and_lines(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "code.md"
            path.write_text("# Heading\n\n```python\nprint('source')\n```\n", encoding="utf-8")
            items, _native, _warnings = parse_markdown(path, {}, {"tables": True})
            self.assertEqual([item["kind"] for item in items], ["heading", "paragraph"])
            self.assertIn("print('source')", items[1]["text"])
            self.assertEqual(items[1]["location"]["line_start"], 3)
            self.assertEqual(items[1]["location"]["line_end"], 5)


if __name__ == "__main__":
    unittest.main()
