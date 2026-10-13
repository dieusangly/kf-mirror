"""Format adapters producing source-located blocks for Task 4."""

import re
from collections import Counter
from pathlib import Path


NUMBERED = re.compile(r"^\s*((?:\d+[°.)]|[IVX]+\.)\s*)")
PDF_DECIMAL_HEADING = re.compile(r"^\s*(\d+(?:\.\d+)*)(?:\.)?\s+")
PDF_ROMAN_HEADING = re.compile(r"^\s*[IVX]+\.\s+", re.I)
PDF_ANNEX_HEADING = re.compile(r"^\s*annexe\s+\d+\b", re.I)
PDF_CASE_HEADING = re.compile(r"^\s*(?:cas|case)\s+n[°o]\s*\d+\b", re.I)
PDF_BOX_HEADING = re.compile(r"^\s*(?:encadré|box)\s+n[°o]\s*\d+\b", re.I)
PICTURE_CONNECTORS = {"a", "à", "au", "aux", "d", "de", "des", "du", "dans", "en",
                      "et", "le", "la", "les", "l", "pour", "par", "sur", "un", "une",
                      "the", "of", "in", "to", "for", "and", "with"}
FIGURE_HEADING = re.compile(r"\b(?:schéma|diagramme|figure|flowchart|diagram)\b", re.I)
URL = re.compile(r"(?:https?://|www\.)\S+", re.I)
PDF_CONTROLS = re.compile(r"[\x00-\x08\x0b-\x1f\x7f]")
WRAPPED_WORD = re.compile(r"(?<=[a-zà-ÿ])-\s+(?=[a-zà-ÿ])", re.I)
HEADER_LABEL = re.compile(r"\s+(?=(?:Les|Le|La|Critères)\s)")
FILENAME_FOOTNOTE = re.compile(
    r"^\s*(?:\d+[.)]?\s+)?[^\s]+\.(?:pdf|docx?|xlsx?|pptx?|html?)\s*$", re.I)


def _pdf_text(value):
    """Remove text-layer controls and join a word broken by a layout hyphen."""
    return WRAPPED_WORD.sub("", PDF_CONTROLS.sub("", value or ""))


def _pdf_table_cells(item):
    """Return table cells with repaired, source-positioned header slots.

    Some narrow PDF tables have adjacent, vertical header labels merged by
    Docling into one cell despite distinct physical columns. Split only those
    labels at a new header phrase, then use their x positions to retain the
    original column mapping for the table body.
    """
    cells = getattr(getattr(item, "data", None), "table_cells", []) or []
    if not cells:
        return []
    raw = sorted(cells, key=lambda cell: (cell.start_row_offset_idx,
                                          cell.start_col_offset_idx))
    header_cells = [cell for cell in raw if getattr(cell, "column_header", False)]
    parts_by_cell = {
        id(cell): [part for part in HEADER_LABEL.split(" ".join(_pdf_text(cell.text).split())) if part]
        if getattr(cell, "column_header", False) else [" ".join(_pdf_text(cell.text).split())]
        for cell in raw
    }
    repaired_headers = any(len(parts_by_cell[id(cell)]) > 1 for cell in header_cells)
    lefts = sorted(getattr(getattr(cell, "bbox", None), "l", None)
                   for cell in header_cells if getattr(getattr(cell, "bbox", None), "l", None) is not None)
    gaps = [right - left for left, right in zip(lefts, lefts[1:]) if 8 <= right - left <= 35]
    column_width = gaps[len(gaps) // 2] if gaps else None
    header_slots = []
    normalized = []
    next_header_column = 0
    for cell in raw:
        row = cell.start_row_offset_idx
        column = cell.start_col_offset_idx
        values = parts_by_cell[id(cell)]
        is_header = getattr(cell, "column_header", False)
        if is_header:
            column = max(column, next_header_column)
            next_header_column = column + len(values)
        elif repaired_headers:
            left = getattr(getattr(cell, "bbox", None), "l", None)
            if left is not None and header_slots:
                column = min(header_slots, key=lambda slot: abs(slot[1] - left))[0]
        for offset, value in enumerate(values):
            slot = column + offset
            if is_header:
                left = getattr(getattr(cell, "bbox", None), "l", None)
                if left is not None:
                    header_slots.append((slot, left + offset * (column_width or 0)))
            normalized.append({
                "row": row,
                "column": slot,
                "text": value,
                "row_span": getattr(cell, "row_span", 1) or 1,
                "col_span": 1 if len(values) > 1 else getattr(cell, "col_span", 1) or 1,
                "header": bool(getattr(cell, "column_header", False) or
                               getattr(cell, "row_header", False)),
                "bbox": getattr(cell, "bbox", None),
            })
    return normalized


def _pdf_table_text(item, doc):
    """Serialize recognized cells once while retaining their column positions."""
    cells = _pdf_table_cells(item)
    if not cells:
        return _pdf_text(item.export_to_markdown(doc=doc)) if hasattr(item, "export_to_markdown") else ""
    rows = {}
    for cell in cells:
        rows.setdefault(cell["row"], []).append(cell)
    lines = []
    for _, row in sorted(rows.items()):
        values = []
        for cell in sorted(row, key=lambda value: value["column"]):
            value = cell["text"]
            column = cell["column"]
            span = cell["col_span"]
            if len(values) <= column:
                values.extend([""] * (column + 1 - len(values)))
            values[column] = value
            if len(values) < column + span:
                values.extend([""] * (column + span - len(values)))
        while values and not values[-1]:
            values.pop()
        if any(values):
            lines.append(" | ".join(values))
    return "\n".join(lines)


def _pdf_running_headers(native):
    """Find text repeated in the top strip of three or more PDF pages."""
    pages = native.get("pages", {})
    matches = {}
    for node in native.get("texts", []):
        value = node.get("text", "")
        for part in node.get("prov") or []:
            page_no = part.get("page_no")
            page = pages.get(str(page_no)) or pages.get(page_no) or {}
            height = page.get("size", {}).get("height")
            box = part.get("bbox")
            if not height or not box:
                continue
            top = box["t"] if box.get("coord_origin") == "TOPLEFT" else height - box["t"]
            if top > height * .16:
                continue
            start, end = part.get("charspan", [0, len(value)])
            fragment = " ".join(_pdf_text(value[start:end]).split())
            if len(fragment) >= 30:
                matches.setdefault(fragment, set()).add(page_no)
    return {fragment for fragment, seen_pages in matches.items() if len(seen_pages) >= 3}


def _picture_contact_furniture(group, pages):
    """Recognize a URL-only publisher footer, not URLs in body prose."""
    value = " ".join(node.get("text", "") for node in group)
    urls = URL.findall(value)
    residual = URL.sub("", value).strip(" \t\n,;|•")
    if len(urls) < 2 or len(re.findall(r"[^\W\d_]+", residual)) > 4:
        return False
    provenances = [node.get("prov") or [] for node in group]
    if not pages or any(len(prov) != 1 for prov in provenances):
        return False
    parts = [prov[0] for prov in provenances]
    page_no = parts[0].get("page_no")
    if page_no != max(int(page) for page in pages) or any(
            part.get("page_no") != page_no for part in parts):
        return False
    page = pages.get(str(page_no)) or pages.get(page_no) or {}
    height = page.get("size", {}).get("height")
    boxes = [part.get("bbox") for part in parts]
    if not height or not all(boxes):
        return False
    return all((min(box["t"], box["b"]) if box.get("coord_origin") != "TOPLEFT" else
                height - max(box["t"], box["b"])) <= height * .2 for box in boxes)


def _picture_group_roles(groups):
    """Separate document flow from text whose context is the figure image."""
    values = [" ".join(node.get("text", "").strip() for node in group) for group in groups]
    flow_page = sum(len(re.findall(r"[^\W\d_]+", value)) >= 12 and len(value) >= 80
                    for value in values) >= 2
    diagram = not flow_page
    roles = []
    for group, value in zip(groups, values):
        label = group[0].get("label", "text")
        if label in {"section_header", "title"} and FIGURE_HEADING.search(value):
            roles.append("flow" if flow_page else "visual")
            diagram = True
            continue
        if not diagram and (label in {"section_header", "title", "list_item",
                                      "checkbox_selected", "checkbox_unselected"} or
                            value.lstrip().startswith(("- ", "• ")) or
                            len(re.findall(r"[^\W\d_]+", value)) >= 9 and len(value) >= 55):
            roles.append("flow")
        elif (diagram and label == "text" and len(value) >= 120 and
              value.rstrip().endswith((".", "?", "!")) and not FIGURE_HEADING.search(value)):
            roles.append("flow")
        else:
            roles.append("visual")
    return roles


def _picture_text_quality(value, label, vocabulary=None):
    """Classify text alone; diagram edges and nearby objects carry no implied meaning."""
    mixed_case = sum(bool(re.search(r"[a-zà-ÿ][A-Z]", word))
                     for word in re.findall(r"[^\W\d_]+", value, flags=re.UNICODE))
    if mixed_case >= 2 and re.search(r"\d{5,}", value):
        return "corrupt_ocr"
    words = re.findall(r"[^\W\d_]+", value, flags=re.UNICODE)
    lowered = {word.casefold() for word in words}
    letters = sum(map(len, words))
    if len(words) < 2 or letters < 6:
        return "graphical"
    if (re.search(r"\d\s*%", value) and len(words) >= 2 and
            any(len(word) >= 4 for word in words) and bool(lowered & PICTURE_CONNECTORS)):
        return "canonical"
    if words[-1].islower() and words[-1].casefold() in PICTURE_CONNECTORS:
        return "potential"
    if label in {"section_header", "title", "caption", "footnote", "list_item",
                 "checkbox_selected", "checkbox_unselected"}:
        return "canonical"
    if len(words) >= 3 and len(words[0]) >= 4 and words[-1].islower() and (
            words[-1].casefold().endswith(("é", "ée", "és", "ées"))):
        return "canonical"
    content_words = [word.casefold() for word in words
                     if len(word) >= 3 and word.casefold() not in PICTURE_CONNECTORS]
    if (vocabulary is not None and 2 <= len(content_words) and letters < 80 and
            not value.rstrip().endswith((";", ":", ".", "?", "!")) and
            not re.search(r"\d\s*(?:%|M€|k€|€)", value, flags=re.I) and
            sum(vocabulary[word] < 2 for word in content_words) >= 2):
        return "potential"
    if (len(words) >= 5 and letters >= 30 or
            len(words) >= 3 and bool(lowered & PICTURE_CONNECTORS) or
            len(words) >= 2 and value.rstrip().endswith((";", ":", ".", "?", "!")) or
            len(words) == 2 and all(len(word) >= 4 for word in words) and
            words[1].islower() and vocabulary is not None and
            all(vocabulary[word.casefold()] >= 2 for word in words)):
        return "canonical"
    return "potential" if (len(words) >= 3 and letters >= 12 or
                           len(words) == 2 and all(len(word) >= 4 for word in words) and
                           words[1].islower()) else "graphical"


def _picture_text_groups(nodes):
    """Join only consecutive fragments on a line or an obvious wrapped line."""
    groups = []
    for node in nodes:
        if not groups:
            groups.append([node])
            continue
        prior = groups[-1][-1]
        prior_prov, node_prov = prior.get("prov") or [], node.get("prov") or []
        a, b = (prior_prov or [{}])[0], (node_prov or [{}])[0]
        box_a, box_b = a.get("bbox"), b.get("bbox")
        same_page = a.get("page_no") == b.get("page_no") and a.get("page_no") is not None
        plain_text = prior.get("label", "text") == node.get("label", "text") == "text"
        text_a, text_b = prior.get("text", "").strip(), node.get("text", "").strip()
        safe_text = (text_a and text_b and not text_a.endswith((";", ":", ".", "?", "!"))
                     and not text_b.startswith(("-", "•", "→"))
                     and not re.fullmatch(r"[\d\W]+", text_b))
        adjacent = False
        if (same_page and plain_text and safe_text and box_a and box_b and
                len(prior_prov) == len(node_prov) == 1):
            height = max(abs(box_a["t"] - box_a["b"]), abs(box_b["t"] - box_b["b"]), 1)
            centre_a = (box_a["t"] + box_a["b"]) / 2
            centre_b = (box_b["t"] + box_b["b"]) / 2
            same_line = (abs(centre_a - centre_b) <= height * .55 and
                         box_b["l"] >= box_a["l"] and box_b["l"] - box_a["r"] <= height * .75)
            overlap = min(box_a["r"], box_b["r"]) - max(box_a["l"], box_b["l"])
            wrapped = (height * .55 < abs(centre_a - centre_b) <= height * 1.6 and
                       overlap >= .4 * min(box_a["r"] - box_a["l"],
                                           box_b["r"] - box_b["l"]) and
                       (text_b[0].islower() or text_a.casefold().split()[-1] in PICTURE_CONNECTORS))
            adjacent = same_line or wrapped
        if adjacent:
            groups[-1].append(node)
        else:
            groups.append([node])
    return groups


def _pdf_toc_levels(native):
    """Use the first document index's cell indentation as heading evidence."""
    indexes = [table for table in native.get("tables", [])
               if table.get("label") == "document_index" and table.get("data", {}).get("table_cells")]
    if not indexes:
        return {}
    table = min(indexes, key=lambda item: (item.get("prov") or [{}])[0].get("page_no", 0))
    cells = table["data"]["table_cells"]
    left_edges = sorted({round(cell["bbox"]["l"], 0) for cell in cells
                         if cell.get("bbox") and cell.get("start_col_offset_idx") == 0})
    if not left_edges:
        return {}
    levels = {}
    index_stack = []
    for cell in sorted(cells, key=lambda item: (item.get("start_row_offset_idx", 0),
                                                 item.get("start_col_offset_idx", 0))):
        if cell.get("start_col_offset_idx") != 0 or not cell.get("bbox"):
            continue
        title = re.sub(r"[_\b]+.*$", "", cell.get("text", "")).strip()
        if title:
            indent = round(cell["bbox"]["l"], 0)
            depth = left_edges.index(indent) + 1
            while index_stack and index_stack[-1][0] >= depth:
                index_stack.pop()
            parent_title = index_stack[-1][1] if index_stack else None
            levels[" ".join(title.casefold().split())] = (depth, parent_title)
            index_stack.append((depth, title))
    return levels


def _pdf_heading_level(text, stack, toc_levels):
    match = PDF_DECIMAL_HEADING.match(text)
    if match and (len(match.group(1).split(".")) > 1 or int(match.group(1)) < 20):
        return len(match.group(1).split(".")), "numbered"
    if PDF_ROMAN_HEADING.match(text) or PDF_ANNEX_HEADING.match(text):
        return 1, "numbered"
    toc_entry = toc_levels.get(" ".join(text.casefold().split()))
    if toc_entry:
        _, parent_title = toc_entry
        parent = next((entry for entry in reversed(stack)
                       if parent_title and " ".join(entry[2].casefold().split()) ==
                       " ".join(parent_title.casefold().split())), None)
        # A missing TOC parent in the body cannot be safely invented.
        return (parent[0] + 1 if parent else 1), "toc"
    anchor = next((entry for entry in reversed(stack) if entry[3] in {"numbered", "toc"}), None)
    case = next((entry for entry in reversed(stack) if entry[3] == "case"), None)
    if PDF_CASE_HEADING.match(text):
        return (anchor[0] if anchor else 1) + 1, "case"
    if PDF_BOX_HEADING.match(text):
        return (anchor[0] if anchor else 1) + 1, "box"
    if case:
        return case[0] + 1, "local"
    if anchor:
        return anchor[0] + 1, "local"
    return 1, "local"


def _item(kind, text, locator_key, location, *, parent_key=None, details=None):
    return {
        "kind": kind,
        "text": text.strip(),
        "locator_key": locator_key,
        "location": location,
        "parent_key": parent_key,
        "details": details or {},
    }


def _css_path(node):
    """Produce an unambiguous CSS path in the saved DOM."""
    parts = []
    while node is not None and isinstance(node.tag, str):
        parent = node.getparent()
        siblings = [child for child in parent] if parent is not None else [node]
        same = [child for child in siblings if child.tag == node.tag]
        part = node.tag
        if len(same) > 1:
            part += f":nth-of-type({same.index(node) + 1})"
        parts.append(part)
        node = parent
    return " > ".join(reversed(parts))


def parse_html(path, source, options):
    from lxml import html

    document = html.parse(str(path))
    root = document.getroot()
    selector = options.get("content_selector", "main")
    matches = root.cssselect(selector)
    if len(matches) != 1:
        raise ValueError(f"HTML content selector {selector!r} matched {len(matches)} nodes; expected one")
    content = matches[0]
    article_id = options.get("article_id") or source["version"]
    items = []
    headings = []
    heading_keys = []

    def add(node, kind, text, parent_key=None, details=None):
        text = " ".join(text.split())
        if not text:
            return None
        css = _css_path(node)
        selected = root.cssselect(css)
        if len(selected) != 1 or selected[0] is not node:
            raise ValueError(f"HTML element selector is not unique: {css}")
        location = {"format": "html", "article_id": article_id, "selector": css}
        match = NUMBERED.match(text)
        if match:
            location["clause"] = match.group(1).strip()
        if node.get("id"):
            location["anchor"] = node.get("id")
        key = f"html:{article_id}:{css}"
        if parent_key is None and heading_keys:
            parent_key = heading_keys[-1]
        items.append(_item(kind, text, key, location, parent_key=parent_key,
                           details={"heading_path": list(headings), **(details or {})}))
        return key

    def visit(node):
        if not isinstance(node.tag, str) or node.tag in {"script", "style", "nav", "aside", "form", "button"}:
            return
        tag = node.tag.lower()
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            level = int(tag[1])
            del headings[level - 1:]
            del heading_keys[level - 1:]
            text = " ".join(node.itertext())
            headings.append(" ".join(text.split()))
            key = add(node, "heading", text, details={"level": level})
            if key:
                heading_keys.append(key)
            return
        if tag == "table":
            rows = []
            for row in node.xpath(".//tr"):
                cells = row.xpath("./th|./td")
                rows.append([" ".join(" ".join(cell.itertext()).split()) for cell in cells])
            table_key = add(node, "table", "\n".join(" | ".join(row) for row in rows),
                            details={"rows": len(rows)})
            for row_index, row in enumerate(node.xpath(".//tr")):
                for cell_index, cell in enumerate(row.xpath("./th|./td")):
                    add(cell, "table_cell", " ".join(cell.itertext()), parent_key=table_key,
                        details={"row": row_index, "column": cell_index,
                                 "header": cell.tag.lower() == "th",
                                 "header_origin": "html_th" if cell.tag.lower() == "th" else "html_td",
                                 "row_span": int(cell.get("rowspan", "1")),
                                 "col_span": int(cell.get("colspan", "1"))})
            return
        if tag == "li":
            text = " ".join(node.xpath("./text() | ./p//text()"))
            add(node, "list_item", text)
            for child in node:
                if child.tag in {"ul", "ol", "table"}:
                    visit(child)
            return
        if tag in {"dt", "dd"}:
            add(node, "list_item", " ".join(node.itertext()))
            return
        if tag == "p":
            add(node, "paragraph", " ".join(node.itertext()))
            return
        if tag == "pre":
            add(node, "paragraph", "".join(node.itertext()), details={"html_type": "pre"})
            return
        if tag == "figcaption":
            add(node, "caption", " ".join(node.itertext()))
            return
        for child in node:
            visit(child)

    if content.tag.lower() == "article":
        # Saved article pages contain controls and related links after the body.
        for heading in content.xpath(".//h1|.//h2"):
            if "article_title" in heading.get("class", "").split():
                visit(heading)
        for body in content.xpath(".//div[contains(concat(' ', normalize-space(@class), ' '), ' content ')]"):
            if "fr-text--md" in body.get("class", "").split():
                visit(body)
    else:
        visit(content)
    if not items:
        raise ValueError("HTML content selector produced no text blocks")
    return items, None, []


def parse_markdown(path, _source, options):
    from markdown_it import MarkdownIt

    text = path.read_text(encoding="utf-8-sig")
    parser = MarkdownIt("commonmark")
    if options.get("tables", True):
        parser.enable("table")
    tokens = parser.parse(text)
    items = []
    heading_path = []
    heading_keys = []
    index = 0
    list_depth = 0
    while index < len(tokens):
        token = tokens[index]
        if token.type == "list_item_open":
            list_depth += 1
        elif token.type == "list_item_close":
            list_depth -= 1
        if token.type in {"heading_open", "paragraph_open", "list_item_open"}:
            # Paragraphs inside list items are emitted by their enclosing list item.
            if token.type == "paragraph_open" and (token.hidden or list_depth):
                index += 1
                continue
            body = tokens[index + 1] if index + 1 < len(tokens) else token
            if token.type == "list_item_open":
                content = []
                depth = 1
                cursor = index + 1
                while cursor < len(tokens) and depth:
                    if tokens[cursor].type == "list_item_open":
                        depth += 1
                    elif tokens[cursor].type == "list_item_close":
                        depth -= 1
                    elif depth == 1 and tokens[cursor].type == "inline":
                        content.append(tokens[cursor].content)
                    cursor += 1
                body_text = " ".join(content)
                if token.map:
                    source_line = text.splitlines()[token.map[0]]
                    marker = re.match(r"^\s*((?:\d+[.)]|[-+*]))\s+", source_line)
                    if marker:
                        body_text = f"{marker.group(1)} {body_text}"
                kind = "list_item"
            else:
                body_text = body.content
                kind = "heading" if token.type == "heading_open" else "paragraph"
            if kind == "heading":
                level = int(token.tag[1])
                del heading_path[level - 1:]
                del heading_keys[level - 1:]
                heading_path.append(body_text)
            if token.map and body_text.strip():
                start, end = token.map
                key = f"markdown:line-{start + 1}:block-{len(items)}"
                items.append(_item(kind, body_text, key,
                                   {"format": "markdown", "heading_path": list(heading_path),
                                    "line_start": start + 1, "line_end": end},
                                   parent_key=heading_keys[-1] if heading_keys else None,
                                   details={"level": int(token.tag[1])} if kind == "heading" else {}))
                if kind == "heading":
                    heading_keys.append(key)
        elif token.type == "table_open" and token.map:
            start, end = token.map
            key = f"markdown:line-{start + 1}:table"
            rows = []
            cells = []
            row = -1
            col = 0
            header = False
            for child in tokens[index + 1:]:
                if child.type == "table_close":
                    break
                if child.type == "thead_open":
                    header = True
                elif child.type == "thead_close":
                    header = False
                elif child.type == "tr_open":
                    row += 1
                    col = 0
                    rows.append([])
                elif child.type in {"th_open", "td_open"}:
                    cell_text = child.content
                    cells.append((row, col, header, cell_text))
                    col += 1
                elif child.type == "inline" and cells:
                    r, c, h, _ = cells[-1]
                    cells[-1] = (r, c, h, child.content)
                    rows[r].append(child.content)
            location = {"format": "markdown", "heading_path": list(heading_path),
                        "line_start": start + 1, "line_end": end}
            items.append(_item("table", "\n".join(" | ".join(r) for r in rows), key,
                               location, parent_key=heading_keys[-1] if heading_keys else None,
                               details={"rows": len(rows)}))
            for row_index, col_index, header, cell_text in cells:
                cell_line = min(end, start + (1 if row_index == 0 else row_index + 2))
                cell_location = {**location, "line_start": cell_line, "line_end": cell_line}
                items.append(_item("table_cell", cell_text,
                                   f"{key}:r{row_index}:c{col_index}", cell_location,
                                   parent_key=key,
                                   details={"row": row_index, "column": col_index,
                                            "header": header, "header_origin": "markdown_header_row",
                                            "row_span": 1, "col_span": 1}))
        elif token.type in {"fence", "code_block", "html_block"} and token.map and token.content.strip():
            start, end = token.map
            key = f"markdown:line-{start + 1}:block-{len(items)}"
            items.append(_item("paragraph", token.content, key,
                               {"format": "markdown", "heading_path": list(heading_path),
                                "line_start": start + 1, "line_end": end},
                               parent_key=heading_keys[-1] if heading_keys else None,
                               details={"markdown_type": token.type, "language": token.info.strip()}))
        index += 1
    return items, None, []


def parse_docx(path, _source, _options):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    items = []
    headings = []
    heading_keys = []
    table_index = 0

    def add_table(table, block_index, parent_key):
        nonlocal table_index
        this_table = table_index
        table_index += 1
        key = f"docx:block-{block_index}:table-{this_table}"
        location = {"format": "docx", "block_index": block_index,
                    "table_index": this_table}
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        items.append(_item("table", "\n".join(" | ".join(row) for row in rows),
                           key, location, parent_key=parent_key,
                           details={"rows": len(rows), "heading_path": list(headings)}))
        cell_anchors = {}
        for row_index, row in enumerate(table.rows):
            seen = set()
            for col_index, cell in enumerate(row.cells):
                cell_identity = cell._tc
                if cell_identity in seen:
                    continue
                seen.add(cell_identity)
                if cell_identity in cell_anchors:
                    items[cell_anchors[cell_identity]]["details"]["row_span"] += 1
                    continue
                cell_location = {**location, "row_index": row_index, "cell_index": col_index}
                cell_key = f"{key}:r{row_index}:c{col_index}"
                items.append(_item("table_cell", cell.text, cell_key,
                                   cell_location, parent_key=key,
                                   details={"row": row_index, "column": col_index,
                                            "header": row_index == 0,
                                            "header_origin": "first_row_inference",
                                            "row_span": 1, "col_span": cell.grid_span,
                                            "paragraphs": [part.text for part in cell.iter_inner_content()
                                                           if isinstance(part, Paragraph)]}))
                cell_anchors[cell_identity] = len(items) - 1
                for part in cell.iter_inner_content():
                    if isinstance(part, Table):
                        add_table(part, block_index, cell_key)

    for block_index, block in enumerate(document.iter_inner_content()):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            style = block.style.name if block.style is not None else ""
            match = re.match(r"Heading\s+([1-6])$", style, flags=re.I)
            kind = "heading" if match else ("list_item" if "list" in style.lower() else "paragraph")
            if match:
                level = int(match.group(1))
                del headings[level - 1:]
                del heading_keys[level - 1:]
                headings.append(text)
            location = {"format": "docx", "block_index": block_index}
            key = f"docx:block-{block_index}"
            items.append(_item(kind, text, key, location,
                               parent_key=heading_keys[-1] if heading_keys else None,
                               details={"style": style, "heading_path": list(headings)}))
            if match:
                heading_keys.append(key)
            continue
        add_table(block, block_index, heading_keys[-1] if heading_keys else None)
    return items, None, []


def build_pdf_converter(options):
    """Create one configured Docling converter for a sequential parse run."""
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions, TesseractCliOcrOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption

    pipeline = PdfPipelineOptions()
    pipeline.do_table_structure = options.get("table_structure", True)
    pipeline.table_structure_options.mode = options.get("table_mode", "accurate")
    if options.get("artifacts_path"):
        pipeline.artifacts_path = Path(options["artifacts_path"])
    ocr = options.get("ocr", {})
    mode = ocr.get("mode", "auto")
    pipeline.do_ocr = mode != "off"
    if pipeline.do_ocr:
        pipeline.ocr_options = TesseractCliOcrOptions(
            mode="full_page" if mode == "force" else "pdf_aware_layout_regions",
            lang=ocr.get("languages", ["fra", "eng"]),
            scale=ocr.get("scale", 3.0),
            psm=ocr.get("psm", 3),
            tesseract_cmd=ocr.get("tesseract_cmd", "tesseract"),
            path=ocr.get("tessdata_path"),
        )
    return DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=pipeline)})


def parse_pdf(path, _source, options, converter=None):
    """Parse one PDF, optionally reusing a converter created for this run."""
    converter = converter or build_pdf_converter(options)
    result = converter.convert(path)
    if str(result.status).lower().split(".")[-1] != "success":
        raise ValueError(f"Docling conversion status: {result.status}")
    doc = result.document
    native = doc.export_to_dict()
    items = []
    warnings = []
    page_counts = {}
    heading_stack = []
    toc_levels = _pdf_toc_levels(native)
    running_headers = _pdf_running_headers(native)
    figures_by_page = {}
    for picture in native.get("pictures", []):
        for part in picture.get("prov") or []:
            if part.get("bbox"):
                figures_by_page.setdefault(part["page_no"], []).append(part["bbox"])
    first_page_body = [node for node in native.get("texts", [])
                       if node.get("parent", {}).get("$ref") == "#/body" and
                       any(part.get("page_no") == 1 for part in node.get("prov") or [])]
    page_text_counts = Counter(part.get("page_no") for node in native.get("texts", [])
                               for part in node.get("prov") or [])

    def large_picture(page, box):
        size = native.get("pages", {}).get(str(page)) or native.get("pages", {}).get(page) or {}
        width, height = size.get("size", {}).get("width"), size.get("size", {}).get("height")
        return bool(width and height and (box["r"] - box["l"]) > width * .5 and
                    abs(box["t"] - box["b"]) > height * .3)

    def picture_annotation(page, location):
        box = location.get("bbox")
        if not box:
            return None
        size = native.get("pages", {}).get(str(page)) or native.get("pages", {}).get(page) or {}
        width, height = size.get("size", {}).get("width"), size.get("size", {}).get("height")
        if not width or not height:
            return None
        for figure_box in figures_by_page.get(page, []):
            if ((figure_box["r"] - figure_box["l"]) <= width * .5 or
                    abs(figure_box["t"] - figure_box["b"]) <= height * .1):
                continue
            bottom = (figure_box["b"] if figure_box.get("coord_origin") == "TOPLEFT"
                      else height - figure_box["b"])
            if 0 <= box[1] - bottom <= 70 and box[0] < figure_box["r"]:
                return figure_box
        return None

    cover_page = len(first_page_body) <= 3 and any(
        large_picture(1, box) for box in figures_by_page.get(1, []))
    native_by_ref = {node["self_ref"]: node for collection in ("texts", "pictures", "tables", "groups")
                     for node in native.get(collection, [])}
    vocabulary = Counter(word.casefold() for node in native.get("texts", [])
                         for word in re.findall(r"[^\W\d_]+", node.get("text", ""), flags=re.UNICODE))
    represented_refs = set()
    excluded_picture_refs = set()

    def add_block(kind, text, ref, parent_ref, provenance, stack, *, evidence_role="authoritative",
                  checkbox_state=None, picture_child=False, minimum_heading_level=1):
        if not text.strip() and kind != "figure":
            return None
        if not provenance:
            warnings.append({"code": "unlocated_pdf_block", "ref": ref, "message": text[:120]})
            return None
        first_key = None
        for index, part in enumerate(provenance):
            page = part["page_no"] if isinstance(part, dict) else part.page_no
            if page not in doc.pages:
                warnings.append({"code": "invalid_pdf_page", "ref": ref, "page": page})
                continue
            if kind == "figure":
                piece = text
            elif len(provenance) == 1:
                piece = text
            else:
                span = part.get("charspan") if isinstance(part, dict) else getattr(part, "charspan", None)
                if not span or len(span) != 2 or not 0 <= span[0] <= span[1] <= len(text):
                    warnings.append({"code": "invalid_pdf_charspan", "ref": ref, "page": page,
                                     "provenance_index": index})
                    continue
                piece = text[span[0]:span[1]]
            piece = _pdf_text(piece)
            if not piece.strip() and kind != "figure":
                if _pdf_text(text).strip():
                    warnings.append({"code": "empty_pdf_charspan", "ref": ref, "page": page,
                                     "provenance_index": index})
                continue
            page_counts[page] = page_counts.get(page, 0) + 1
            key = f"pdf:p{page}:block-{page_counts[page]}"
            if first_key is None:
                first_key = key
            numbered = NUMBERED.match(piece)
            location = {"format": "pdf", "page": page,
                        "label": numbered.group(1).strip() if numbered else f"block {page_counts[page]}"}
            bbox = part.get("bbox") if isinstance(part, dict) else getattr(part, "bbox", None)
            if bbox is not None:
                if isinstance(bbox, dict):
                    if bbox.get("coord_origin") == "TOPLEFT":
                        location["bbox"] = [bbox["l"], bbox["t"], bbox["r"], bbox["b"]]
                    else:
                        height = doc.pages[page].size.height
                        location["bbox"] = [bbox["l"], height - bbox["t"],
                                            bbox["r"], height - bbox["b"]]
                else:
                    box = bbox.to_top_left_origin(doc.pages[page].size.height)
                    location["bbox"] = [box.l, box.t, box.r, box.b]
            local_role = evidence_role
            annotation_figure = None
            if kind == "paragraph" and local_role == "authoritative":
                if " ".join(piece.split()) in running_headers:
                    local_role = "structural"
                elif cover_page and page == 1 and len(piece.split()) < 15:
                    local_role = "auxiliary"
                elif len(piece.split()) <= 32 and not piece.rstrip().endswith((".", ";", ":", "?", "!")):
                    annotation_figure = picture_annotation(page, location)
                    if annotation_figure is not None:
                        local_role = "auxiliary"
                if (local_role == "authoritative" and len(piece.strip()) == 1 and
                        piece.strip().isalpha() and location.get("bbox")):
                    box = location["bbox"]
                    size = native.get("pages", {}).get(str(page)) or native.get("pages", {}).get(page) or {}
                    width, height = size.get("size", {}).get("width"), size.get("size", {}).get("height")
                    if width and height and box[0] < width * .2 and box[1] < height * .16:
                        local_role = "structural"
                if (local_role == "authoritative" and page_text_counts[page] <= 10 and
                        len(re.sub(r"\W", "", piece)) <= 4 and location.get("bbox")):
                    size = native.get("pages", {}).get(str(page)) or native.get("pages", {}).get(page) or {}
                    height = size.get("size", {}).get("height")
                    if height and location["bbox"][1] > height * .6:
                        local_role = "structural"
            if kind == "heading":
                heading_level, role = _pdf_heading_level(piece, stack, toc_levels)
                heading_level = max(heading_level, minimum_heading_level)
                while stack and stack[-1][0] >= heading_level:
                    stack.pop()
            else:
                heading_level, role = None, None
            parent_key = (next((item["locator_key"] for item in reversed(items)
                                if item["details"].get("docling_ref") == parent_ref and
                                item["location"]["page"] == page), None) or
                          next((item["locator_key"] for item in reversed(items)
                                if item["details"].get("docling_ref") == parent_ref), None)
                          if picture_child else None)
            if parent_key is None:
                parent_key = stack[-1][1] if stack else None
            if annotation_figure is not None:
                figure_parent = next((item["locator_key"] for item in reversed(items)
                                      if item["kind"] == "figure" and
                                      item["location"]["page"] == page), None)
                parent_key = figure_parent or parent_key
            heading_path = [entry[2] for entry in stack] + ([piece] if kind == "heading" else [])
            details = {"docling_ref": ref, "docling_parent_ref": parent_ref,
                       "level": heading_level, "heading_path": heading_path,
                       "evidence_role": local_role,
                       "provenance": "ocr_or_pdf_text" if pipeline.do_ocr else "pdf_text"}
            if checkbox_state:
                details["checkbox_state"] = checkbox_state
            if annotation_figure is not None:
                details["visual_annotation"] = True
            items.append(_item(kind, piece, key, location, parent_key=parent_key, details=details))
            if kind == "heading":
                stack.append((heading_level, key, piece, role))
        if first_key:
            represented_refs.add(ref)
        return first_key

    def picture_descendants(picture_ref, local_stack):
        picture = native_by_ref.get(picture_ref, {})
        excluded = []
        potential = []
        potential_candidates = []
        visual_uncertain = []
        corrupt = []
        furniture = []
        text_nodes = []
        child_heading_level = local_stack[-1][0] + 1 if local_stack else 1

        def visit(ref):
            node = native_by_ref.get(ref)
            if node is None:
                warnings.append({"code": "unresolved_picture_child", "ref": ref,
                                 "picture_ref": picture_ref})
                return
            if ref.startswith("#/groups/"):
                for child in node.get("children", []):
                    visit(child["$ref"])
                return
            if ref.startswith("#/texts/"):
                if node.get("content_layer") == "furniture" or ref in represented_refs:
                    return
                text_nodes.append(node)
                return
            warnings.append({"code": "unsupported_picture_child", "ref": ref,
                             "picture_ref": picture_ref})

        for child in picture.get("children", []):
            visit(child["$ref"])
        groups = _picture_text_groups(text_nodes)
        for group, role in zip(groups, _picture_group_roles(groups)):
            refs = [node["self_ref"] for node in group]
            value = " ".join(node.get("text", "").strip() for node in group)
            label = group[0].get("label", "text")
            if _picture_contact_furniture(group, native.get("pages", {})):
                furniture.extend(refs)
                excluded_picture_refs.update(refs)
                continue
            quality = _picture_text_quality(value, label, vocabulary)
            if quality != "canonical":
                target = (corrupt if quality == "corrupt_ocr" else
                          visual_uncertain if quality == "potential" and role == "visual" else
                          potential if quality == "potential" else excluded)
                target.extend(refs)
                if quality == "potential":
                    potential_candidates.append({"refs": refs, "text": value,
                                                 "reason": "text_not_reliably_canonical",
                                                 "role": role})
                excluded_picture_refs.update(refs)
                continue
            if len(group) == 1:
                node = group[0]
                provenance = node.get("prov") or []
            else:
                parts = [(node.get("prov") or []) for node in group]
                pages = {part[0].get("page_no") for part in parts if len(part) == 1}
                boxes = [part[0].get("bbox") for part in parts if len(part) == 1]
                if len(pages) != 1 or len(boxes) != len(group) or not all(boxes):
                    (visual_uncertain if role == "visual" else potential).extend(refs)
                    potential_candidates.append({"refs": refs, "text": value,
                                                 "reason": "incomplete_fragment_provenance",
                                                 "role": role})
                    excluded_picture_refs.update(refs)
                    continue
                origin = boxes[0].get("coord_origin", "BOTTOMLEFT")
                if any(box.get("coord_origin", "BOTTOMLEFT") != origin for box in boxes):
                    (visual_uncertain if role == "visual" else potential).extend(refs)
                    potential_candidates.append({"refs": refs, "text": value,
                                                 "reason": "mixed_bbox_origins", "role": role})
                    excluded_picture_refs.update(refs)
                    continue
                provenance = [{"page_no": next(iter(pages)), "bbox": {
                    "l": min(box["l"] for box in boxes), "r": max(box["r"] for box in boxes),
                    "t": max(box["t"] for box in boxes) if origin != "TOPLEFT" else min(box["t"] for box in boxes),
                    "b": min(box["b"] for box in boxes) if origin != "TOPLEFT" else max(box["b"] for box in boxes),
                    "coord_origin": origin}}]
            ref = refs[0]
            kind = ("heading" if label in {"section_header", "title"} else
                    "list_item" if label in {"list_item", "checkbox_selected",
                                             "checkbox_unselected"} else
                    "caption" if label == "caption" else
                    "footnote" if label == "footnote" else "paragraph")
            if role == "flow" and kind == "paragraph" and value.lstrip().startswith(("- ", "• ")):
                kind = "list_item"
            marker = group[0].get("marker")
            if kind == "list_item" and len(group) == 1:
                value = group[0].get("orig") or (f"{marker} {value}" if marker else value)
            first_key = add_block(kind, value, ref, picture_ref, provenance, local_stack,
                                  evidence_role="auxiliary" if role == "visual" else "authoritative",
                                  picture_child=role == "visual",
                                  minimum_heading_level=child_heading_level,
                                  checkbox_state=("selected" if label == "checkbox_selected" else "unselected")
                                  if label in {"checkbox_selected", "checkbox_unselected"} else None)
            if first_key and len(group) > 1:
                represented_refs.update(refs)
                element = next(item for item in reversed(items) if item["locator_key"] == first_key)
                element["details"]["docling_refs"] = refs
                element["details"]["source_parts"] = [
                    {"ref": node["self_ref"], "text": node.get("text", ""),
                     "provenance": node.get("prov") or []} for node in group]
        if excluded:
            warnings.append({"code": "excluded_picture_ocr_fragments", "picture_ref": picture_ref,
                             "count": len(excluded), "refs": excluded})
        if corrupt:
            warnings.append({"code": "excluded_picture_corrupt_ocr", "picture_ref": picture_ref,
                             "count": len(corrupt), "refs": corrupt})
        if furniture:
            warnings.append({"code": "excluded_picture_furniture", "picture_ref": picture_ref,
                             "count": len(furniture), "refs": furniture})
        if visual_uncertain:
            warnings.append({"code": "excluded_figure_text_uncertain", "picture_ref": picture_ref,
                             "count": len(visual_uncertain), "refs": visual_uncertain,
                             "candidates": [candidate for candidate in potential_candidates
                                            if candidate["role"] == "visual"]})
        if potential:
            warnings.append({"code": "potential_picture_text", "picture_ref": picture_ref,
                             "count": len(potential), "refs": potential,
                             "candidates": [candidate for candidate in potential_candidates
                                            if candidate["role"] == "flow"]})

    for item, level in doc.iterate_items(with_groups=False):
        label = str(getattr(item, "label", "")).lower()
        if "section_header" in label or "title" in label:
            kind = "heading"
        elif "list_item" in label:
            kind = "list_item"
        elif "table" in label or label == "document_index":
            kind = "table"
        elif "picture" in label:
            kind = "figure"
        elif "caption" in label:
            kind = "caption"
        elif "footnote" in label:
            kind = "footnote"
        elif label in {"checkbox_selected", "checkbox_unselected"}:
            kind = "list_item"
        else:
            kind = "paragraph"
        if kind == "table":
            text = _pdf_table_text(item, doc)
        else:
            text = getattr(item, "text", "") or ""
        marker = getattr(item, "marker", None)
        if kind == "list_item":
            text = getattr(item, "orig", None) or (f"{marker} {text}" if marker else text)
        if kind == "heading" and re.match(r"^\d{2,}\.\s+.*:\s*$", text):
            kind = "paragraph"  # A numbered quotation lead is not a section.
        if kind == "heading" and text.rstrip().endswith("…"):
            kind = "paragraph"  # A sentence introducing examples is not a section.
        if kind in {"heading", "list_item", "paragraph"} and \
                " ".join(_pdf_text(text).split()) in running_headers:
            kind = "paragraph"  # Repeated top-strip text is publisher furniture.
        if kind == "heading" and len(text.split()) <= 12 and (prov := getattr(item, "prov", None)):
            first = prov[0]
            bbox = getattr(first, "bbox", None)
            if bbox is not None:
                page = first.page_no
                box = bbox.to_top_left_origin(doc.pages[page].size.height)
                if picture_annotation(page, {"bbox": [box.l, box.t, box.r, box.b]}):
                    kind = "paragraph"  # A short label below a large image annotates it.
        if kind == "paragraph" and re.match(r"^[a-z]\.\s+", text, re.I):
            kind = "list_item"  # Alphabetic alternatives in a quoted rule.
        prov = getattr(item, "prov", []) or []
        ref = getattr(item, "self_ref", None)
        parent_ref = getattr(getattr(item, "parent", None), "cref", None)
        if kind == "table" and not text.strip():
            cells = _pdf_table_cells(item)
            if cells:
                rows = {}
                for cell in cells:
                    rows.setdefault(cell["row"], []).append(cell)
                text = "\n".join(" | ".join(cell["text"] for cell in sorted(row,
                    key=lambda value: value["column"])) for _, row in sorted(rows.items()))
            if not text.strip():
                warnings.append({"code": "empty_docling_table", "ref": ref,
                                 "message": "table has no text or cells"})
                continue
        if kind == "table" and len(prov) > 1:
            warnings.append({"code": "multi_page_table", "ref": ref,
                             "message": "table omitted until its cells can be assigned to pages"})
            continue
        filename_only_footnote = kind == "footnote" and FILENAME_FOOTNOTE.match(text)
        key = add_block(kind, text, ref, parent_ref, prov, heading_stack,
                        evidence_role="structural" if kind == "figure" or filename_only_footnote else "authoritative",
                        checkbox_state=("selected" if label == "checkbox_selected" else "unselected")
                        if label in {"checkbox_selected", "checkbox_unselected"} else None)
        if key is None:
            continue
        items[-1]["details"]["docling_group_depth"] = level
        if filename_only_footnote:
            items[-1]["details"]["link_text_role"] = "non_substantive_filename"
        if kind == "figure":
            picture_descendants(ref, list(heading_stack))
        if kind == "table" and getattr(item, "data", None) is not None:
            for provenance in prov:
                page = provenance.page_no
                location = items[-1]["location"]
                for cell in _pdf_table_cells(item):
                    row = cell["row"]
                    col = cell["column"]
                    cell_location = dict(location)
                    cell_bbox = cell["bbox"]
                    if cell_bbox is not None:
                        try:
                            box = cell_bbox.to_top_left_origin(doc.pages[page].size.height)
                            cell_location["bbox"] = [box.l, box.t, box.r, box.b]
                        except (AttributeError, TypeError, KeyError):
                            cell_location.pop("bbox", None)
                    items.append(_item("table_cell", cell["text"],
                                       f"{key}:r{row}:c{col}", cell_location,
                                       parent_key=key,
                                       details={"row": row, "column": col,
                                                "row_span": cell["row_span"],
                                                "col_span": cell["col_span"],
                                                "header": cell["header"],
                                                "header_origin": "docling_table_structure",
                                                "evidence_role": "structural"}))
    for node in native.get("texts", []):
        ref = node["self_ref"]
        if (node.get("content_layer") != "furniture" and ref not in represented_refs and
                ref not in excluded_picture_refs and _pdf_text(node.get("text", "")).strip()):
            warnings.append({"code": "unrepresented_pdf_text", "ref": ref,
                             "parent_ref": node.get("parent", {}).get("$ref"),
                             "message": node.get("text", "")[:120]})
    # A multi-page Docling node can be yielded before the previous page's
    # footnotes. Restore page order while retaining Docling's within-page order.
    items.sort(key=lambda item: item["location"]["page"])
    return items, native, warnings
